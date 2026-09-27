"""Produce revision 1.4 of the Dynamic Apps Standard SRS from revision 1.3 (built by update_srs_v13.py).

Adds standards and requirements management: a Standards Catalog (internal frameworks and external standards,
structured from chapter to requirement), a Requirements Catalog (all requirement origins, typed), and the
applicability of standards and requirements to an Organization and/or a project, with compliance tracking.
Formats are cloned from the document itself; the manual TOC page numbers are set afterwards by srs_toc_pages.py.

usage: python3 update_srs_v14.py <srs_v1.3.docx> <srs_v1.4.docx>
"""
import copy
import sys

import docx
from docx.oxml.ns import qn

SRC, OUT = sys.argv[1], sys.argv[2]
d = docx.Document(SRC)
W = qn('w:t')


def para(start, style=None):
    for p in d.paragraphs:
        if p.text.startswith(start) and (style is None or p.style.name.startswith(style)):
            return p
    raise KeyError(start)


def set_runs(p, texts):
    runs = p._p.findall(qn('w:r'))
    for i, r in enumerate(runs):
        if i < len(texts):
            ts = r.findall(W)
            for t in ts[1:]:
                r.remove(t)
            ts[0].text = texts[i]
            ts[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        else:
            p._p.remove(r)
    return p


class Cursor:
    def __init__(self, anchor):
        self.el = anchor._p if hasattr(anchor, '_p') else anchor

    def add(self, template, texts):
        el = copy.deepcopy(template._p)
        self.el.addnext(el)
        self.el = el
        return set_runs(docx.text.paragraph.Paragraph(el, template._parent), texts)

    def req(self, rid, text): return self.add(T_REQ, [f'{rid}: ', text])
    def body(self, text): return self.add(T_BODY, [text])
    def bullet(self, text): return self.add(T_BUL, [text])
    def h2(self, text): return self.add(T_H2, [text])
    def h1(self, text): return self.add(T_H1, [text])


def last_req(prefix):
    return [p for p in d.paragraphs if p.text.startswith(prefix)][-1]


def add_rows(table, rows):
    for cells in rows:
        tr = copy.deepcopy(table.rows[-2]._tr)
        table.rows[-1]._tr.addnext(tr)
        for tc, text in zip(tr.findall(qn('w:tc')), cells):
            ps = tc.findall(qn('w:p'))
            for extra in ps[1:]:
                tc.remove(extra)
            set_runs(docx.text.paragraph.Paragraph(ps[0], None), [text])


T_REQ = para('FR-DA-TEN-01')
T_BUL = para('Explainable AI only')
T_BODY = para('Every Dynamic App is built on')
T_H2 = para('3.10 Return on Experience', 'Heading')
T_H1 = para('Appendix B', 'Heading')
T_TOC2 = para('4.35 Gate Library', 'Normal')
T_TOC1 = para('Appendix E — Traceability to the Verticals and SME Requirements', 'Normal')

T_REF = para('CortexPLM — Reference implementation')


def retext(p, old, new):
    for t in p._p.iter(W):
        if t.text and old in t.text:
            t.text = t.text.replace(old, new)
            return p
    raise KeyError(old)


# ---------------------------------------------------------------- cover, purpose, references, overview
set_runs(para('Version 1.3'), ['Version 1.4  ·  September 2026  ·  Confidential'])
Cursor(para('Revision 1.3 integrates the application-agnostic')).body(
    'Revision 1.4 adds standards and requirements management for any Dynamic App: a Standards Catalog holding internal '
    'frameworks and external standards (ISO, IEC, regulations, industry and customer standards) structured from part and '
    'chapter down to requirement; a Requirements Catalog in which every requirement, whatever its origin, is one typed '
    'record; and the applicability of a whole standard or of individual requirements to an Organization and/or a project, '
    'with tailoring, compliance status, evidence and a statement of applicability. Appendix F lists the changes.')
c = Cursor(para('Application-Agnostic Requirements for Verticals and SME —'))
c.add(T_REF, ['ISO/IEC/IEEE 29148:2018 — Systems and software engineering — Life cycle processes — Requirements engineering.'])
c.add(T_REF, ['OMG Requirements Interchange Format (ReqIF), Version 1.2.'])
Cursor(para('Revision 1.3 adds Section 3.14')).body(
    'Revision 1.4 adds Section 3.15, Sections 4.36 to 4.38, new requirements in Sections 4.3, 4.5, 4.6, 4.23, 4.27, 5.1, '
    '5.4, 5.9 and 5.10, and Appendix F, the change log of revision 1.4. Nothing was removed or renumbered.')

# ---------------------------------------------------------------- 1.4 definitions
add_rows(d.tables[0], [
    ('Standard', 'A body of requirements kept in the Standards Catalog, either Internal (an Organization\'s own framework, policy or method) or External (an international, national or industry standard, a regulation, or a customer standard), identified by code, title, issuer and edition.'),
    ('Standard Structure', 'The ordered tree of a standard — part, chapter, clause, sub-clause to any depth — whose leaves are requirements.'),
    ('Requirement', 'A single, verifiable statement kept once in the Requirements Catalog, with a type, an obligation level (shall, should, may), a verification method and an owner; it may belong to a standard or stand alone.'),
    ('Requirement Type', 'A configurable classification of requirements by origin or nature, for example Standard, Internal, Regulatory, Contractual, Product, Process, Security & Privacy, Quality.'),
    ('Applicability', 'The link that makes a standard, a clause or a requirement apply to an Organization and/or a project, with its status: Applicable, Not applicable (justified) or Tailored.'),
    ('Statement of Applicability', 'The list of every requirement in scope for an Organization or a project, with its applicability, justification, compliance status and evidence.'),
    ('Compliance Status', 'The assessed state of an applicable requirement: Not assessed, Compliant, Partially compliant, Non-compliant, or Not applicable.'),
])

# ---------------------------------------------------------------- 2.3 product functions
Cursor(para('Activate verticals and an SME mode')).bullet(
    'Manage internal and external standards from chapter to requirement, keep every requirement in one typed catalog, '
    'apply standards and requirements to Organizations and projects, and track compliance with evidence.')

# ---------------------------------------------------------------- 3.15 architecture
anchor = d.paragraphs[[i for i, p in enumerate(d.paragraphs) if p.text.startswith('4. Common Functional Requirements') and p.style.name.startswith('Heading')][0] - 1]
c = Cursor(anchor)
c.h2('3.15 Standards & Requirements Architecture')
c.body('Standards and requirements management is built on three catalogs with one rule: a requirement exists once. The '
       'Standards Catalog holds each standard and its structure (part > chapter > clause > sub-clause, to any depth). The '
       'Requirements Catalog holds every requirement as a typed record; a standard requirement is a catalog requirement '
       'placed at a leaf of its standard\'s structure, and an internal, regulatory, contractual or product requirement is '
       'the same kind of record, with or without a standard. The Applicability register links standards, clauses or '
       'single requirements to an Organization and/or a project and carries the compliance assessment.')
c.body('Applicability is inherited downward: what applies to an Organization applies to all its projects, and a project '
       'adds its own standards or requirements and may tailor inherited ones only with a justification; a mandatory '
       'Organization requirement cannot be set to Not applicable in a project without an approved waiver. Requirements '
       'are linked to each other (derives from, satisfies, equivalent to, conflicts with) so that one piece of evidence '
       'can prove several standards, and they are linked to the objects that implement or verify them: Controls, '
       'checklist items, gate criteria, processes, tasks and attachments. The Compliance & Security Standards modules of '
       'Section 4.3 and the vertical standards of Section 4.27 are entries of the same Standards Catalog, and the catalogs '
       'are indexed by the retrieval engine of Section 3.5.')
c.body('Two ownership levels keep content governed and private: platform entries (external standards, reference '
       'frameworks) are published once and shared read-only with every Organization; Organization entries (internal '
       'frameworks, customer requirements) stay private to their tenant and may be shared read-only with the other '
       'Organizations of the same Group. Where a standard\'s text is copyrighted, the catalog stores the clause reference, '
       'title and a paraphrase, and the full text only under a licence recorded with the standard.')

# ---------------------------------------------------------------- links in existing sections
Cursor(last_req('FR-DA-CFG-')).req('FR-DA-CFG-15', 'The system shall record every Compliance & Security Standards module of FR-DA-CFG-04 as an entry of the Standards Catalog (Section 4.36), so that activating the module also applies the standard to the Organization (Section 4.38) and seeds its Controls from the standard\'s requirements (FR-DA-CFG-05).')
Cursor(last_req('FR-DA-KB-')).req('FR-DA-KB-05', 'The system shall index the Standards Catalog and the Requirements Catalog with the knowledge base, so that AI grounding and search return the clause and requirement that apply, only from standards the user may see.')
Cursor(last_req('FR-DA-GOV-')).req('FR-DA-GOV-09', 'The system shall let a Business Rule, a Control or a Risk/Opportunity reference one or more catalog requirements (Section 4.37), refining the standard-level tag of FR-DA-GOV-03 down to the requirement.')
Cursor(last_req('FR-DA-CHK-')).req('FR-DA-CHK-08', 'The system shall let a checklist item and a gate criterion reference one or more catalog requirements, so that completing the item or passing the gate records evidence for those requirements (Section 4.38).')
Cursor(last_req('FR-DA-VCP-')).req('FR-DA-VCP-04', 'The system shall hold vertical regulations and standards in the Standards Catalog and link them to their vertical, so that activating a vertical proposes them for applicability to the Organization.')

# ---------------------------------------------------------------- 4.36 – 4.38
c = Cursor(last_req('FR-DA-GTE-'))
c.h2('4.36 Standards Catalog')
c.body('The Standards Catalog holds every standard the Organization works with, internal or external, structured from chapter to requirement.')
c.req('FR-DA-STD-01', 'The system shall keep a Standards Catalog with full Create, Read, Update, Delete and Duplicate on standards, each with a code, title, category (Internal or External), kind (international or national standard, regulation, industry standard, customer standard, internal framework, policy, method), issuer, edition or version, publication and effective dates, language, owner, licence note, and applicable sectors or verticals.')
c.req('FR-DA-STD-02', 'The system shall let an authorized user build the structure of a standard as an ordered tree of parts, chapters, clauses and sub-clauses to any depth, each with a number, title and optional guidance, and place requirements of the Requirements Catalog (Section 4.37) at its leaves; nodes can be added, moved, renumbered and removed while the standard is in Draft.')
c.req('FR-DA-STD-03', 'The system shall manage each standard through Draft, Published, Superseded and Withdrawn; only published standards shall be applicable, and a new edition shall be created from the previous one, with a comparison listing added, changed and removed requirements.')
c.req('FR-DA-STD-04', 'The system shall import a standard\'s structure and requirements from a spreadsheet (CSV or XLSX, one row per node) or ReqIF, validating numbering and parent references and reporting every rejected row, and export any standard in the same formats and to PDF.')
c.req('FR-DA-STD-05', 'The system shall let an authorized user draft a standard\'s structure and requirements from an uploaded document with AI, as an Assistive suggestion reviewed node by node (FR-DA-AI-07), and shall record the source document and page of every requirement so drafted.')
c.req('FR-DA-STD-06', 'The system shall publish platform standards read-only to every Organization, keep Organization standards private to their tenant with optional read-only sharing inside the Group, restrict catalog changes to a dedicated permission, and audit every change.')
c.req('FR-DA-STD-07', 'The system shall present the catalog with search and filters (category, kind, issuer, sector, status, applied or not) and, for each standard, its structure as a collapsible tree with requirement counts per node and the Organizations and projects it applies to.')

c.h2('4.37 Requirements Catalog')
c.body('Every requirement — from a standard, an internal framework, a regulation, a contract or a product specification — is one typed record, stored once and reused everywhere.')
c.req('FR-DA-RQM-01', 'The system shall keep a Requirements Catalog with full Create, Read, Update, Delete and Duplicate on requirements, each with a unique identifier, title, statement, type, obligation level (Mandatory "shall", Recommended "should", Optional "may"), rationale, verification method (inspection, analysis, demonstration, test, audit, document review), expected evidence, source (standard and clause, document or person), owner, tags, status and version.')
c.req('FR-DA-RQM-02', 'The system shall let an administrator manage the list of requirement types with CRUD, seeded with Standard, Internal, Regulatory, Contractual, Product, Process, Security & Privacy and Quality; a type in use shall be deactivated rather than deleted.')
c.req('FR-DA-RQM-03', 'The system shall let a requirement exist with or without a standard; a requirement placed in a standard\'s structure shows its clause path, and the same requirement shall not be duplicated when it is reused in another standard or project.')
c.req('FR-DA-RQM-04', 'The system shall let a requirement be linked to other requirements as derives from, refines, satisfies, equivalent to or conflicts with, and show the links both ways, so that an internal requirement can be traced to the external clauses it covers.')
c.req('FR-DA-RQM-05', 'The system shall show for every requirement where it is used: standards and clauses, Organizations and projects it applies to, and the Controls, checklist items, gate criteria, processes, tasks and attachments linked to it.')
c.req('FR-DA-RQM-06', 'The system shall manage each requirement through Draft, Approved and Retired, version every change under Section 3.8, and never alter an assessment already recorded against an earlier version.')
c.req('FR-DA-RQM-07', 'The system shall present the catalog with search and filters (type, obligation, source standard, status, owner, tag, applied or not), bulk edit of type, owner and tags, and CSV, XLSX and ReqIF import and export.')
c.req('FR-DA-RQM-08', 'The system shall suggest with AI, for review, duplicate or equivalent requirements and links between requirements of different standards; no link shall be saved without a person accepting it.')

c.h2('4.38 Applicability & Compliance')
c.body('Standards and requirements apply to an Organization, to a project, or to both; compliance is assessed where they apply.')
c.req('FR-DA-APL-01', 'The system shall let an authorized user apply a whole standard, a clause with all requirements below it, or single requirements to an Organization, to one or more projects, or to both, and shall show on the Organization and on each project the standards and requirements that apply.')
c.req('FR-DA-APL-02', 'The system shall make requirements applied to an Organization apply to all its projects; a project may add standards and requirements, and may mark an inherited requirement Tailored or Not applicable only with a justification, and, for a mandatory one, an approved waiver.')
c.req('FR-DA-APL-03', 'The system shall let a project template (Section 4.33) and a vertical (Section 4.25) carry default standards and requirements, applied to a project created from them.')
c.req('FR-DA-APL-04', 'The system shall record, for every applicable requirement in an Organization or a project, its compliance status, owner, due date, evidence (attachments, Controls, completed checklist items, gate decisions), last assessment date and assessor, and shall keep the history of assessments.')
c.req('FR-DA-APL-05', 'The system shall produce, for an Organization or a project, a Statement of Applicability and a compliance matrix (requirements × status) with coverage per standard and per clause, the list of gaps, and export to PDF, XLSX and CSV (Section 4.8).')
c.req('FR-DA-APL-06', 'The system shall alert the owner through the Notification Center when a requirement becomes Non-compliant, when an assessment is overdue, and when a new edition of an applied standard is published, and shall list the requirements affected by the new edition.')
c.req('FR-DA-APL-07', 'The system shall record the application, tailoring, waiver and assessment of every requirement in the audit trail, and shall never present a compliance status as a certification (FR-DA-CFG-09).')

# ---------------------------------------------------------------- NFR
Cursor(last_req('NFR-DA-PERF-')).req('NFR-DA-PERF-09', 'Importing a standard of 5,000 requirements shall complete within 1 minute, and the compliance matrix of a project with 2,000 applicable requirements shall display within 3 seconds.')
Cursor(last_req('NFR-DA-SCALE-')).req('NFR-DA-SCALE-06', 'The catalogs shall support at least 500 standards and 100,000 requirements per tenant without functional degradation.')
Cursor(last_req('NFR-DA-COMPAT-')).req('NFR-DA-COMPAT-04', 'Standards and requirements shall be exchanged in CSV, XLSX and ReqIF 1.2, verified by a round-trip import and export that preserves structure, identifiers, types and links.')
c = Cursor(last_req('NFR-DA-COMP-'))
c.req('NFR-DA-COMP-07', 'The system shall store the full text of a copyrighted standard only under a licence recorded with it, and otherwise only its clause references, titles and paraphrased requirements.')
c.req('NFR-DA-COMP-08', 'Requirements in the catalog shall follow the characteristics of ISO/IEC/IEEE 29148 (necessary, unambiguous, verifiable, singular), and the system shall flag on entry a requirement without an obligation level or verification method.')

# ---------------------------------------------------------------- 6 data model
add_rows(d.tables[2], [
    ('Standard (catalog)', 'A Standards Catalog entry: code, title, category (Internal or External), kind, issuer, edition, dates, status, owner, licence, ownership level (platform or Organization).'),
    ('StandardNode', 'A node of a standard\'s structure — part, chapter, clause or sub-clause — with number, title, guidance, parent and order.'),
    ('Requirement', 'A Requirements Catalog entry: identifier, title, statement, type, obligation, verification method, expected evidence, source, owner, tags, status, version; optionally placed at a StandardNode.'),
    ('RequirementType', 'A configurable requirement type with its active flag.'),
    ('RequirementLink', 'A typed link between two requirements (derives from, refines, satisfies, equivalent to, conflicts with), or between a requirement and a Control, checklist item, gate criterion, process, task or attachment.'),
    ('Applicability', 'A standard, node or requirement applied to an Organization and/or a project, with status, justification, waiver and inheritance source.'),
    ('ComplianceAssessment', 'A dated assessment of an applicable requirement: status, evidence, assessor, comments.'),
])

# ---------------------------------------------------------------- 7.1 / 7.2
Cursor(para('Revision 1.3 adds screens for verticals')).body(
    'Revision 1.4 adds a Standards Catalog (list and structure tree), a Requirements Catalog with types and links, a '
    'Standards & Requirements tab on each Organization and project, and a compliance matrix with its Statement of '
    'Applicability.')
c = Cursor(para('Packs and onboarding: pack and bundle CRUD'))
c.bullet('Standards: catalog CRUD, duplicate, lifecycle and editions with comparison; structure tree CRUD and reorder; import and export (CSV, XLSX, ReqIF, PDF); AI draft from a document.')
c.bullet('Requirements: catalog CRUD, bulk edit, versions; requirement types CRUD; links between requirements and to implementing objects; where-used; import and export; AI duplicate and link suggestions.')
c.bullet('Applicability and compliance: apply and unapply to an Organization or a project; tailoring and waivers; assessments with evidence; Statement of Applicability and compliance matrix export.')

# ---------------------------------------------------------------- Appendix A row
add_rows(d.tables[3], [
    ('Standards & requirements management (Sections 4.36 – 4.38)', 'Standards Catalog from chapter to requirement; typed Requirements Catalog; applicability to Organizations and projects; compliance tracking and Statement of Applicability.',
     'New in revision 1.4. Target. CortexPLM covers part of it: GDPR, ISO/IEC 27001 and SOC 2 are toggled per Organization and seed tagged Controls, and a reference knowledge library supports AI grounding. The catalogs, the structure tree, requirement types and links, and project-level applicability are not implemented.'),
])

# ---------------------------------------------------------------- Appendix F (change log)
closing = para('Product-specific material — the content of each vertical instance')
c = Cursor(closing)
c.h1('Appendix F — Revision 1.4 Change Log')
c.body('Revision 1.4 is additive: every earlier requirement keeps its ID and wording, and new requirements take the next free number in their section or a new family.')
tbl = copy.deepcopy(d.tables[6]._tbl)
c.el.addnext(tbl)
F = docx.table.Table(tbl, closing._parent)
for tr in F.rows[3:]:
    tbl.remove(tr._tr)
CHANGES = [
    ('Cover, 1.1, 1.5, 1.6', 'Version 1.4', 'Version raised to 1.4; purpose and overview describe the revision; ISO/IEC/IEEE 29148 and ReqIF added to the references.'),
    ('1.4 Definitions', '7 terms added', 'Standard, Standard Structure, Requirement, Requirement Type, Applicability, Statement of Applicability, Compliance Status.'),
    ('2.3 Product Functions', '1 function added', 'Standards and requirements management with applicability and compliance.'),
    ('3.15 (new)', 'Architecture', 'Three catalogs, one record per requirement, inherited applicability, requirement links, ownership levels, licensed text.'),
    ('4.3 Configuration Management', 'FR-DA-CFG-15', 'Compliance & Security Standards modules are Standards Catalog entries.'),
    ('4.5 Knowledge Base', 'FR-DA-KB-05', 'Catalogs indexed for AI grounding and search.'),
    ('4.6 Governance', 'FR-DA-GOV-09', 'Business Rules, Controls and Risks reference catalog requirements.'),
    ('4.23 Checklist Library', 'FR-DA-CHK-08', 'Checklist items and gate criteria reference requirements and record evidence.'),
    ('4.27 Vertical Compliance', 'FR-DA-VCP-04', 'Vertical standards held in the catalog and proposed on vertical activation.'),
    ('4.36 Standards Catalog', 'FR-DA-STD-01 – 07 (new)', 'Internal and external standards with CRUD; structure from part and chapter to requirement; lifecycle and editions; import and export; AI draft; sharing and permissions; browsing.'),
    ('4.37 Requirements Catalog', 'FR-DA-RQM-01 – 08 (new)', 'Typed requirements with CRUD; configurable types; one record reused; links between requirements; where-used; lifecycle and versions; filters, bulk edit, exchange; AI suggestions.'),
    ('4.38 Applicability & Compliance', 'FR-DA-APL-01 – 07 (new)', 'Apply to Organization and/or project; inheritance and tailoring; defaults from templates and verticals; assessments with evidence; Statement of Applicability and matrix; alerts; audit.'),
    ('5.1, 5.4, 5.9, 5.10', 'NFR-DA-PERF-09, SCALE-06, COMPAT-04, COMP-07 – 08', 'Import and matrix performance; catalog scale; CSV, XLSX and ReqIF round-trip; licensed text; requirement quality per ISO/IEC/IEEE 29148.'),
    ('6 Data Model', '7 entities', 'Standard (catalog), StandardNode, Requirement, RequirementType, RequirementLink, Applicability, ComplianceAssessment.'),
    ('7.1, 7.2 Interfaces', 'Clarified', 'New screens and endpoint families for standards, requirements, applicability and compliance.'),
    ('Appendix A', '1 row', 'Conformance of the 1.4 additions in CortexPLM.'),
]
for r, cells in zip(F.rows[1:3], CHANGES[:2]):
    for cell, text in zip(r.cells, cells):
        set_runs(cell.paragraphs[0], [text])
add_rows(F, CHANGES[2:])
c.el = tbl

# ---------------------------------------------------------------- manual TOC entries
Cursor(para('3.14 Vertical & SME Layer Architecture', 'Normal')).add(T_TOC2, ['3.15 Standards & Requirements Architecture', '0'])
c = Cursor(T_TOC2)
for t in ('4.36 Standards Catalog', '4.37 Requirements Catalog', '4.38 Applicability & Compliance'):
    c.add(T_TOC2, [t, '0'])
Cursor(T_TOC1).add(T_TOC1, ['Appendix F — Revision 1.4 Change Log', '0'])

d.save(OUT)
print('saved', OUT)
