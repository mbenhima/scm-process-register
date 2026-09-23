"""Shared Word styling for the CortexPLM guides (A4, Times New Roman, orange/grey palette, TOC, page X of Y)."""
import os
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches, Twips

ORANGE, ORANGE_DEEP, ORANGE_TINT = 'F8931D', 'E07B00', 'FDEEDA'
GREY_DARK, GREY_INK, GREY_MED, GREY_LIGHT, GREY_LINE = '3A3A3C', '58595B', '808184', 'F2F2F3', 'E3E3E4'
FONT = 'Times New Roman'
CONTENT_WIDTH = 6.77  # inches: A4 width 8.27 minus 2 x 0.75 margins
LOGO = os.path.join(os.path.dirname(__file__), '..', 'brand', 'poweract-logo.png')

rgb = lambda h: RGBColor.from_string(h)

ORDER = {
    'pPr': ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl', 'numPr', 'suppressLineNumbers', 'pBdr', 'shd', 'tabs',
            'suppressAutoHyphens', 'kinsoku', 'wordWrap', 'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi', 'adjustRightInd',
            'snapToGrid', 'spacing', 'ind', 'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc', 'textDirection', 'textAlignment',
            'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr', 'sectPr', 'pPrChange'],
    'tblPr': ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize', 'tblW', 'jc', 'tblCellSpacing', 'tblInd',
              'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption', 'tblDescription'],
    'tcPr': ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar', 'textDirection', 'tcFitText', 'vAlign', 'hideMark'],
}


def put(parent, el):
    """Insert el into a *Pr element at its schema position, replacing an element with the same tag."""
    kind = parent.tag.split('}')[1]; order = ORDER[kind]; name = el.tag.split('}')[1]
    old = parent.find(el.tag)
    if old is not None: parent.remove(old)
    pos = order.index(name)
    for i, child in enumerate(list(parent)):
        cname = child.tag.split('}')[1]
        if cname in order and order.index(cname) > pos:
            parent.insert(i, el); return el
    parent.append(el); return el



def _font(run, size=None, bold=None, italic=None, color=None, name=FONT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run._element.rPr.rFonts.set(qn('w:cs'), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color: run.font.color.rgb = rgb(color)
    return run


def _border(el_pr, side, color, sz=8, space=4):
    bdr = el_pr.find(qn('w:pBdr'))
    if bdr is None:
        bdr = put(el_pr, OxmlElement('w:pBdr'))
    b = OxmlElement(f'w:{side}')
    for k, v in (('val', 'single'), ('sz', str(sz)), ('space', str(space)), ('color', color)): b.set(qn(f'w:{k}'), v)
    bdr.append(b)


def new_document(title):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(sec, m, Twips(1080))
    st = doc.styles
    normal = st['Normal']
    normal.font.name, normal.font.size = FONT, Pt(12)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    normal.font.color.rgb = rgb(GREY_DARK)
    pf = normal.paragraph_format
    pf.line_spacing, pf.space_after, pf.alignment = 1.15, Pt(6), WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size, color, before, after in (('Heading 1', 17, ORANGE_DEEP, 18, 8), ('Heading 2', 13.5, GREY_DARK, 14, 6), ('Heading 3', 12, ORANGE_DEEP, 10, 4)):
        h = st[name]
        h.font.name, h.font.size, h.font.bold, h.font.italic = FONT, Pt(size), True, False
        h.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        h.element.rPr.rFonts.set(qn('w:ascii'), FONT)
        h.element.rPr.rFonts.set(qn('w:hAnsi'), FONT)
        h.font.color.rgb = rgb(color)
        rf = h.element.rPr.rFonts
        for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(before), Pt(after)
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h.paragraph_format.keep_with_next = True
        if name == 'Heading 1':
            _border(h.element.get_or_add_pPr(), 'bottom', ORANGE_DEEP, sz=8, space=4)
    for lst in ('List Bullet', 'List Number'):
        s = st[lst]; s.font.name, s.font.size = FONT, Pt(12)
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.core_properties.title = title
    doc.core_properties.author = 'POWERACT Consulting'
    return doc


def _field(paragraph, instr, placeholder=''):
    r = paragraph.add_run(); fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'begin'); r._r.append(fc)
    r2 = paragraph.add_run(); it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = f' {instr} '; r2._r.append(it)
    r3 = paragraph.add_run(); fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'separate'); r3._r.append(fc)
    r4 = paragraph.add_run(placeholder)
    r5 = paragraph.add_run(); fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'end'); r5._r.append(fc)
    for r in (r, r2, r3, r4, r5): _font(r, size=9, color=GREY_MED)
    return paragraph


def header_footer(doc, credit):
    for sec in doc.sections:
        sec.different_first_page_header_footer = True
        hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _font(hp.add_run(credit), size=9, color=GREY_MED)
        _border(hp._p.get_or_add_pPr(), 'bottom', GREY_LINE, sz=4, space=4)
        fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(fp.add_run('POWERACT Consulting  ·  Page '), size=9, color=GREY_MED)
        _field(fp, 'PAGE', '1'); _font(fp.add_run(' of '), size=9, color=GREY_MED); _field(fp, 'NUMPAGES', '1')


def cover(doc, eyebrow, title, subtitle, meta):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(90); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(LOGO):
        p.add_run().add_picture(LOGO, width=Inches(2.4))
    else:  # text wordmark in the brand colours when the logo file is not available
        _font(p.add_run('POWER'), size=30, bold=True, color=GREY_DARK)
        _font(p.add_run('ACT'), size=30, bold=True, color=ORANGE)
        p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(p2.add_run('C O N S U L T I N G'), size=11, bold=True, color=GREY_INK)
    sp = doc.add_paragraph(); sp.paragraph_format.space_before = Pt(80); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(sp.add_run(eyebrow.upper()), size=11, bold=True, color=ORANGE_DEEP)
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(t.add_run(title), size=30, bold=True, color=GREY_DARK)
    s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(s.add_run(subtitle), size=14, color=GREY_INK)
    m = doc.add_paragraph(); m.paragraph_format.space_before = Pt(120); m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(m.add_run(meta), size=11, color=GREY_MED)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def toc(doc):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _font(p.add_run('Table of Contents'), size=17, bold=True, color=ORANGE_DEEP)
    _border(p._p.get_or_add_pPr(), 'bottom', ORANGE_DEEP, sz=8)
    tp = doc.add_paragraph()
    _field(tp, 'TOC \\o "1-3" \\h \\z \\u', 'Right-click and choose Update Field to build the table of contents.')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def h1(doc, text, new_page=True):
    if new_page and len(doc.paragraphs) > 3:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    return doc.add_heading(text, 1)


def para(doc, text='', bold_prefix=None, italic=False, align=None, size=None, color=None):
    p = doc.add_paragraph()
    if bold_prefix: _font(p.add_run(bold_prefix), bold=True, size=size)
    _rich(p, text, size=size, italic=italic, color=color)
    if align == 'left': p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p


def _rich(p, text, size=None, italic=False, color=None):
    """Supports **bold** spans (used for what the user types or clicks)."""
    parts = text.split('**')
    for i, part in enumerate(parts):
        if part: _font(p.add_run(part), size=size, bold=(i % 2 == 1) or None, italic=italic or None, color=color)
    return p


def bullets(doc, items, style='List Bullet'):
    for it in items:
        p = doc.add_paragraph(style=style); _rich(p, it)
        p.paragraph_format.space_after = Pt(2)


def numbered(doc, items):
    """Numbered steps rendered as '1.' text so numbering restarts in every procedure."""
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.35); p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(3)
        _font(p.add_run(f'{i}.  '), bold=True, color=ORANGE_DEEP)
        _rich(p, it)


def _cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor); put(tcPr, shd)


def _table_borders(table, color=GREY_LINE):
    tblPr = table._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), color); b.append(e)
    put(tblPr, b)


def _fixed_layout(t, widths):
    tblPr = t._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); put(tblPr, lay)
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(sum(widths) * 1440))); tw.set(qn('w:type'), 'dxa'); put(tblPr, tw)
    for i, col in enumerate(t.columns): col.width = Inches(widths[i])
    grid = t._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn('w:gridCol'))): gc.set(qn('w:w'), str(int(widths[i] * 1440)))


def table(doc, headers, rows, widths=None, size=9.5, status_col=None):
    widths = widths or [CONTENT_WIDTH / len(headers)] * len(headers)
    scale = min(1, CONTENT_WIDTH / sum(widths)); widths = [w * scale for w in widths]
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    _table_borders(t)
    _fixed_layout(t, widths)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; _cell_bg(c, ORANGE); c.width = Inches(widths[i])
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT; _font(p.add_run(h), size=size, bold=True, color='FFFFFF')
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); trPr.append(th)
    for r_i, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            c = cells[i]; c.width = Inches(widths[i])
            if r_i % 2 == 1: _cell_bg(c, GREY_LIGHT)
            p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            lines = str(val).split('\n')
            for li, line in enumerate(lines):
                if li: p = c.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
                _rich(p, line, size=size)
            if status_col is not None and i == status_col:
                _cell_bg(c, {'Covered': 'D9EAD3', 'Met': 'D9EAD3', 'Fixed': 'B6D7A8', 'Partial': 'FBE0B5', 'Deployment': 'FFF3B0', 'Missing': 'F4C7C3', 'Gap': 'F4C7C3'}.get(str(val).replace('*', '').split(' ')[0], 'D9EAD3'))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def callout(doc, text, title=None):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _fixed_layout(t, [CONTENT_WIDTH]); c = t.rows[0].cells[0]; _cell_bg(c, ORANGE_TINT); c.width = Inches(CONTENT_WIDTH)
    tblPr = t._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for side in ('top', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), 'nil'); b.append(e)
    e = OxmlElement('w:left'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '24'); e.set(qn('w:color'), ORANGE); b.insert(1, e)
    put(tblPr, b)
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if title: _font(p.add_run(title + '  '), bold=True, size=11, color=ORANGE_DEEP)
    _rich(p, text, size=11)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figure(doc, path, caption, width=CONTENT_WIDTH):
    if not os.path.exists(path): return
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Inches(width))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(c.add_run(caption), size=9.5, italic=True, color=GREY_INK)
