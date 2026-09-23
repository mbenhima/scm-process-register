"""Produce revision 1.1 of the Dynamic Apps Standard SRS from the 1.0 source.

Edits a copy of the source document in place, cloning its own paragraph, heading,
TOC and table formats so new content matches the original layout. The page numbers
in the manual table of contents are set afterwards by set_toc_pages() from the PDF.

usage: python3 update_srs.py <source.docx> <out.docx>
"""
import copy
import re
import sys

import docx
from docx.oxml.ns import qn

SRC, OUT = sys.argv[1], sys.argv[2]
d = docx.Document(SRC)
W = qn('w:t')


def para(start, style=None):
    for p in d.paragraphs:
        if p.text.startswith(start) and (style is None or p.style.name.startswith(style)):
            return p
    raise KeyError(start)


def set_runs(p, texts):
    """Put texts into the paragraph's first len(texts) runs, dropping the rest."""
    runs = p._p.findall(qn('w:r'))
    for i, r in enumerate(runs):
        if i < len(texts):
            ts = r.findall(W)
            for t in ts[1:]:
                r.remove(t)
            ts[0].text = texts[i]
            ts[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        else:
            p._p.remove(r)
    return p


def clone_after(anchor_el, template_p, texts):
    el = copy.deepcopy(template_p._p)
    anchor_el.addnext(el)
    p = docx.text.paragraph.Paragraph(el, template_p._parent)
    set_runs(p, texts)
    return p


# Templates taken from the source document itself.
T_REQ = para('FR-DA-TEN-01')                       # bold ID run + text run
T_BUL = para('Explainable AI only')                # plain bullet
T_BODY = para('Every Dynamic App is built on')     # body paragraph
T_H2 = para('3.10 Return on Experience', 'Heading')
T_H1 = para('Appendix A', 'Heading')
T_TOC2 = para('3.10 Return on Experience', 'Normal')
T_TOC1 = [p for p in d.paragraphs if p.text.startswith('Appendix A') and p.style.name == 'Normal'][0]
T_REF = para('journi — Software Requirements')


class Cursor:
    """Inserts a run of new blocks after an anchor element, in order."""
    def __init__(self, anchor):
        self.el = anchor._p if hasattr(anchor, '_p') else anchor

    def add(self, template, texts):
        p = clone_after(self.el, template, texts)
        self.el = p._p
        return p

    def req(self, rid, text): return self.add(T_REQ, [f'{rid}: ', text])
    def bullet(self, text): return self.add(T_BUL, [text])
    def body(self, text): return self.add(T_BODY, [text])
    def h2(self, text): return self.add(T_H2, [text])
    def h1(self, text): return self.add(T_H1, [text])


def last_req(prefix):
    return [p for p in d.paragraphs if p.text.startswith(prefix)][-1]


def add_rows(table, rows):
    for cells in rows:
        src = table.rows[-2]._tr        # keeps the alternating row shading
        tr = copy.deepcopy(src)
        table.rows[-1]._tr.addnext(tr)
        for tc, text in zip(tr.findall(qn('w:tc')), cells):
            ps = tc.findall(qn('w:p'))
            for extra in ps[1:]:
                tc.remove(extra)
            set_runs(docx.text.paragraph.Paragraph(ps[0], None), [text])


def append_cell(table, row_start, text):
    for r in table.rows:
        if r.cells[0].text.startswith(row_start):
            p = r.cells[-1].paragraphs[-1]
            runs = p._p.findall(qn('w:r'))
            last = runs[-1].findall(W)[-1]
            last.text = last.text + text
            return
    raise KeyError(row_start)


# ---------------------------------------------------------------- cover
set_runs(para('Version 1.0'), ['Version 1.1  ·  September 2026  ·  Confidential'])
cons = para('Consolidated from journi')
set_runs(cons, ['Consolidated from journi (DynamicJourni) v2.0 and NCP Solver v2.0, extended from the CortexPLM reference build'])

# ---------------------------------------------------------------- 1.1 Purpose
Cursor(para('This document is intended to serve')).body(
    'Revision 1.1 extends the standard with platform capabilities first built and verified in CortexPLM, an '
    'innovation-lifecycle Dynamic App, where they proved independent of that application\'s business domain — '
    'for example the pinnable, repositionable navigation menu, internal benchmarking, and in-place upgrade and '
    'backup. Only requirements applicable to any Dynamic App were retained; CortexPLM\'s phase-gate lifecycle and '
    'its product-specific content remain in its own documentation. Appendix B lists every change made in this revision.')

# ---------------------------------------------------------------- 1.2 Scope
Cursor(para('In scope:')).body(
    'Revision 1.1 adds to this scope: a personalizable navigation shell (a menu pinned open or sliding away, docked '
    'on any edge of the screen, with favorites that follow the user); internal benchmarking within an Organization '
    'and across the Organizations of one Group; and platform operations (scheduled backup with retention, additive '
    'startup migrations, startup guards, a health endpoint, a reproducible demonstration seed, and continuous '
    'integration across operating systems).')
Cursor(para('Out of scope:')).body(
    'Also out of scope: external benchmarking — any comparison with an organization outside the requester\'s Group, '
    'or with third-party or industry data.')

# ---------------------------------------------------------------- 1.4 Definitions
add_rows(d.tables[0], [
    ('Navigation Shell', 'The frame shared by every authenticated screen: the primary menu, the header bar (context switcher, language, notifications, AI Assistant, user menu), and the content area.'),
    ('Pinned / Sliding Menu', 'The two menu modes. Pinned, the menu stays open and the content reflows beside it. Unpinned, it collapses to an edge handle and slides over the content when pointed at, focused, or toggled.'),
    ('Dock Position', 'The screen edge — left, right, top, or bottom — a user chooses for the primary menu; a per-user preference stored server-side.'),
    ('Internal Benchmarking', 'Comparison of standard metrics between segments of one Organization, or between Organizations of the same Group, using aggregates only. External benchmarking is out of scope.'),
    ('Minimum Sample Size', 'The smallest number of records a benchmark segment or Organization must contain to be ranked; smaller samples are shown as not comparable.'),
    ('Additive Migration', 'A change applied at startup that only adds tables, columns, permission codes, or default grants, and never drops or rewrites existing data or an administrator\'s edits.'),
])

# ---------------------------------------------------------------- 1.5 References
Cursor(para('Dynamic Apps Suite — Brand')).add(T_REF, [
    'CortexPLM — Reference implementation of the Dynamic Apps platform layer, build of September 2026, with its SRS Gap Analysis, September 2026.'])

# ---------------------------------------------------------------- 1.6 Overview
Cursor(para('Section 2 describes the Dynamic Apps Suite')).body(
    'Revision 1.1 adds Sections 3.11 to 3.13 and 4.19 to 4.21, adds requirements to existing sections under new, '
    'never-reused IDs, and adds Appendix B, the change log of this revision. No 1.0 requirement was removed or renumbered.')

# ---------------------------------------------------------------- 2.2 Design principles
c = Cursor(para('Multilingual and RTL-first'))
c.bullet('One navigation shell, personal to each user. Every Dynamic App shares one navigation shell; each user may pin '
         'the menu or let it slide away, dock it on any screen edge, and mark favorites, and these choices follow the user across sessions and devices.')
c.bullet('Compare within the family, never outside it. Benchmarking compares an Organization with itself and with its own '
         'Group, from aggregates only; it never exposes another tenant\'s records and never reaches beyond the Group.')
c.bullet('Upgrade and recover in place. A new release upgrades an existing installation at startup without data loss, and '
         'the data is backed up on a schedule the operator does not have to remember.')

# ---------------------------------------------------------------- 2.3 Functions
c = Cursor(para('Provide a searchable, multi-language in-app Help'))
c.bullet('Provide a personalizable navigation shell: menu pinned or sliding, docked left, right, top, or bottom, with favorites, persisted per user.')
c.bullet('Provide internal benchmarking of standard metrics between segments of one Organization and between Organizations of the same Group.')
c.bullet('Protect and upgrade data in place: on-demand and daily backups with retention, and additive migrations at startup.')

# ---------------------------------------------------------------- 2.5 Operating environment
set_runs(para('Server: a standard Node.js runtime'), [
    'Server: a standard Node.js runtime (Node.js 18+ minimum; Node.js 22.13+ recommended — the first long-term-support '
    'line whose built-in SQLite module works without an experimental flag — so the platform can use the runtime\'s own '
    'SQLite module rather than a native-compiled dependency). The required version shall be declared in the package manifest.'])

# ---------------------------------------------------------------- 2.6 Constraints / 2.7 Assumptions
Cursor(para('Every user-facing string must exist')).bullet(
    'Benchmark data leaving an Organization, even towards its own Group, must be aggregated; no record, name, or identifier of another tenant may be returned.')
Cursor(para('Demo/seed data is fictional')).bullet(
    'Off-site copies of backups, and the hosting platform\'s own disaster-recovery arrangements, remain the responsibility of the operator of each deployment.')

# ---------------------------------------------------------------- 3.1 Tenancy
Cursor(para('Every Dynamic App is built on the same three')).body(
    'Group membership is also the only boundary across which Organizations may compare performance (Section 3.12). '
    'Every hierarchy screen shows whether an Organization belongs to a Group (Yes/No) and which one; an Organization '
    'with no Group is shown as Independent.')

# ---------------------------------------------------------------- 3.11 – 3.13
c = Cursor(para('Any closable entity'))
c.h2('3.11 Navigation Shell & Personalization Architecture')
c.body('Every authenticated screen renders inside one shared navigation shell: a primary menu, a header bar carrying the '
       'Organization/Project context switcher, language switcher, notification indicator, AI Assistant entry point and '
       'user menu, and the content area. The menu is generated from a single navigation model in which each item declares '
       'its module group, route, icon, and the view permission and Solution Pack entitlement it requires — the same codes '
       'the server enforces on the route — so an item a user cannot open is never shown, and a new module appears in '
       'navigation by being registered once.')
c.body('Each user chooses where the menu is docked (left, right, top, or bottom) and whether it is pinned open or slides '
       'away. Unpinned, the menu collapses to an edge handle and slides over the content when pointed at, focused, or '
       'toggled. These choices, the user\'s favorite items, and collapsed menu groups are stored server-side as user '
       'preferences, so they follow the user to any browser or device. Left and right are logical sides: under Arabic RTL '
       'they mirror. Below a mobile breakpoint the shell switches to a drawer, whatever the stored dock position.')
c.h2('3.12 Internal Benchmarking Architecture')
c.body('Internal benchmarking computes a catalog of comparable metrics — each defined once with its unit, its direction of '
       'good (higher or lower is better), and its computation from the platform\'s own records, including the audit trail '
       'where history matters — and aggregates them by segment. Two scopes exist. Within an Organization, segments are '
       'dimensions of the Organization\'s own data (for example record or project type, process track, or owning '
       'department), and authorized users keep full drill-down to the underlying records. Across a Group, each '
       'participating Organization contributes only per-metric aggregates; no record, name, or identifier leaves its tenant.')
c.body('Every segment or Organization below the minimum sample size is flagged as not comparable and excluded from ranks '
       'and medians. An Organization Administrator can withdraw the Organization from Group sharing at any time, after '
       'which its figures are no longer shown to the rest of the Group. Comparison with organizations outside the Group, '
       'or with external or third-party data, is out of scope.')
c.h2('3.13 Platform Operations Architecture')
c.body('Operational robustness belongs to the platform layer, not to each deployment. On every start the server applies '
       'additive migrations — creating missing tables and columns and adding new permission codes with their default role '
       'grants — without dropping data or overwriting an administrator\'s edits, so an existing installation upgrades in '
       'place. Startup guards report an unsupported runtime, a port already in use, or an empty database, each as a '
       'plain-language message naming the command that fixes it, rather than a stack trace.')
c.body('The persistence layer takes consistent online backups on demand and on a daily schedule, keeps them for a '
       'configurable retention period, and lists them to the platform administrator. A short-lived, per-user read cache, '
       'invalidated by any write in the same Organization, keeps read latency within target under concurrent load without '
       'ever serving one user\'s or one tenant\'s data to another.')

# ---------------------------------------------------------------- 4 intro
Cursor(para('Each requirement below is common to journi')).body(
    'Requirements added in revision 1.1 are generalized from the CortexPLM reference build; they carry new IDs '
    'following the last 1.0 ID of their section, are listed in Appendix B, and their conformance status is recorded in Appendix A.')

# ---------------------------------------------------------------- 4.1
c = Cursor(last_req('FR-DA-TEN-'))
c.req('FR-DA-TEN-09', 'The system shall display, on every Organization list and detail screen, whether the Organization belongs to a Group (Yes/No) and, if so, which Group; an Organization with no Group shall be labeled Independent.')
c.req('FR-DA-TEN-10', 'The system shall reject, server-side, an OBS parent assignment that crosses trees — an Organization-level node under a Project-level node, or a node under another Project\'s tree — and an assignment that makes a node its own ancestor.')
c.req('FR-DA-TEN-11', 'The system shall identify records by values that reveal nothing about other tenants\' activity — opaque identifiers (e.g. UUIDs) or per-tenant sequences — and shall answer a request for another tenant\'s record exactly as for a record that does not exist (HTTP 404), in support of FR-DA-TEN-08.')

# ---------------------------------------------------------------- 4.2
c = Cursor(last_req('FR-DA-RBAC-'))
c.req('FR-DA-RBAC-08', 'The system shall, when a user opens a screen or deep link that their roles or the Organization\'s configuration do not allow, show an in-app access notice explaining why and whom to ask, never a blank page or a raw error.')
c.req('FR-DA-RBAC-09', 'The system shall provide a dedicated permission for viewing the content of individual evaluation or verification verdicts; other users, except the parties to that verification, shall see that an evaluation exists but not its content.')
c.req('FR-DA-RBAC-10', 'The system shall, at startup, add any permission code introduced by a new release to the catalog together with its default role grants, without overwriting grants an administrator has changed (Section 3.13).')

# ---------------------------------------------------------------- 4.11
Cursor(last_req('FR-DA-AUD-')).req('FR-DA-AUD-05',
    'The system shall allow an authorized user to reopen a completed workflow step before the stage or approval it belongs to is decided, only with a justification recorded per Section 3.7; the reopening shall return the step to its active state and be written to the audit log.')

# ---------------------------------------------------------------- 4.14
c = Cursor(last_req('FR-DA-I18N-'))
c.req('FR-DA-I18N-07', 'The system shall return every server-generated, user-facing message — validation and error messages, notification text, AI Assistant answers, and export labels — in the requesting user\'s display language, resolved from the same translation dictionary as the UI; a machine code may accompany the translated text but shall not replace it.')
c.req('FR-DA-I18N-08', 'The system shall seed its reference and Knowledge Base content in every supported language and shall let the user filter Knowledge Base articles by language, defaulting to the user\'s display language.')

# ---------------------------------------------------------------- 4.16
Cursor(last_req('FR-DA-VER-')).req('FR-DA-VER-06',
    'The system shall provide one shared version-comparison view, available from the history of every versioned entity, that shows two selected versions side by side and highlights the fields that differ, in support of FR-DA-VER-03.')

# ---------------------------------------------------------------- 4.18
Cursor(last_req('FR-DA-REX-')).req('FR-DA-REX-08',
    'The system shall present the REX prompt to the user who performed the closing action, immediately after the closure succeeds, with a choice to fill it in now or skip; skipping shall leave the closure intact and the entry available to add later from the REX Register.')

# ---------------------------------------------------------------- 4.19 – 4.21
c = Cursor(last_req('FR-DA-REX-'))
c.h2('4.19 Navigation Shell & Personalization')
c.body('The navigation shell (Section 3.11) behaves the same way in every Dynamic App, so a user who has learned one application can find their way around any other.')
c.req('FR-DA-NAV-01', 'The system shall generate the primary menu from a single navigation model in which each item declares its module group, label, icon, route, and the view permission and Solution Pack entitlement it requires, and shall show only the items the current user can open.')
c.req('FR-DA-NAV-02', 'The system shall group menu items by module, let the user collapse and expand each group, and mark the current screen as the active item.')
c.req('FR-DA-NAV-03', 'The system shall let each user pin the menu open or unpin it. Unpinned, the menu shall collapse to an edge handle, slide over the content when the user points at the edge, focuses the handle, or uses the menu button, and slide away again when the pointer leaves or an item is chosen.')
c.req('FR-DA-NAV-04', 'The system shall let each user dock the menu on the left, right, top, or bottom edge of the screen, laying it out vertically on the left and right and horizontally, with drop-down groups, on the top and bottom.')
c.req('FR-DA-NAV-05', 'The system shall let each user mark menu items as favorites, shown in a Favorites group at the start of the menu.')
c.req('FR-DA-NAV-06', 'The system shall store the dock position, pinned state, collapsed groups, and favorites as server-side user preferences, applied on every sign-in and on any device, and shall offer them both from the menu itself and from the user\'s Settings screen.')
c.req('FR-DA-NAV-07', 'The system shall treat left and right as logical start and end under right-to-left languages, mirroring the dock position and the slide direction when Arabic is selected.')
c.req('FR-DA-NAV-08', 'The system shall, below a mobile breakpoint, replace the docked menu with a drawer opened from a menu button and closed by Escape, by an outside tap, or by choosing an item, without changing the user\'s stored desktop preferences.')
c.req('FR-DA-NAV-09', 'The system shall make the shell fully keyboard-operable: a skip-to-content link as the first focusable element, menu and group buttons exposing their expanded state, Escape closing any open menu, and the active item marked as the current page for assistive technologies.')
c.req('FR-DA-NAV-10', 'The system shall keep the Organization/Project context switcher, language switcher, notification indicator, AI Assistant entry point, and user menu in the header bar of every authenticated screen.')

c.h2('4.20 Internal Benchmarking')
c.body('Internal benchmarking (Section 3.12) lets an Organization learn from its own segments and from the other Organizations of its Group. External benchmarking is out of scope.')
c.req('FR-DA-BMK-01', 'The system shall define its benchmark metrics in one catalog module, each with a label, a unit, a direction of good, a plain-language explanation, and a computation from the platform\'s own records, consumed by every benchmarking screen.')
c.req('FR-DA-BMK-02', 'The system shall compare, within one Organization, every catalog metric across the segments of at least one dimension the application defines (for example record or project type, process track, or owning OBS node), showing each segment\'s value, the Organization-wide value, and the median.')
c.req('FR-DA-BMK-03', 'The system shall compare, across the Organizations of the requester\'s Group, per-Organization aggregates of the same metrics, and shall rank each Organization per metric according to the metric\'s direction of good, giving tied values the same rank.')
c.req('FR-DA-BMK-04', 'The system shall include in a Group comparison only per-metric aggregates and each contributing Organization\'s name and sector — never a record, record title, person\'s name, or identifier from another Organization.')
c.req('FR-DA-BMK-05', 'The system shall flag any segment or Organization with fewer records than a minimum sample size (default 3) as not comparable, and shall exclude it from ranks and medians.')
c.req('FR-DA-BMK-06', 'The system shall let an Organization Administrator withdraw the Organization from Group benchmarking and rejoin it; while withdrawn, its figures shall not be shown to the other Organizations of the Group. Every change shall be written to the audit log.')
c.req('FR-DA-BMK-07', 'The system shall gate internal benchmarking behind a view capability, Group comparison behind a separate capability, and the sharing setting behind a management capability, each enforced server-side.')
c.req('FR-DA-BMK-08', 'The system shall refuse any comparison with an Organization outside the requester\'s Group and shall not import or display external or third-party benchmark data; an Organization with no Group shall see its internal comparison only, with a notice explaining why.')
c.req('FR-DA-BMK-09', 'The system shall render each metric comparison as a chart following the common visual identity, with the requester\'s own Organization or the selected segment highlighted, and shall let the user filter the comparison by the application\'s segment dimensions.')

c.h2('4.21 Platform Operations')
c.body('These requirements keep a Dynamic App safe to upgrade, recover, and demonstrate, whatever its business domain (Section 3.13).')
c.req('FR-DA-OPS-01', 'The system shall take a consistent backup of its data store on demand — from the command line and from the platform administrator\'s configuration screen — and automatically once a day, while remaining online.')
c.req('FR-DA-OPS-02', 'The system shall keep backups for a configurable retention period (default 14 days), delete older ones automatically, and list existing backups with date and size to the platform administrator only.')
c.req('FR-DA-OPS-03', 'The system shall document how to restore a backup, and a restored backup from an earlier release shall start without any manual migration step (FR-DA-OPS-04).')
c.req('FR-DA-OPS-04', 'The system shall apply additive schema and permission-catalog migrations automatically on startup, never dropping or rewriting existing data or administrator changes, so an existing installation upgrades in place.')
c.req('FR-DA-OPS-05', 'The system shall check, when installed and started, the runtime version, port availability, and whether the data store has been initialized, and shall report each problem as a plain-language message naming the corrective action rather than a stack trace.')
c.req('FR-DA-OPS-06', 'The system shall expose an unauthenticated health endpoint reporting only liveness, whether the data store is initialized, and the deployment mode — never tenant data.')
c.req('FR-DA-OPS-07', 'The system shall ship a deterministic demonstration seed, reset by one command, containing at least one Group of two or more Organizations, at least one Independent Organization, several sectors, a user for every role, and at least ten instances of each principal end-to-end process per Organization, all fictional (Section 2.7).')
c.req('FR-DA-OPS-08', 'The system shall include an administrator traceability view mapping each requirement ID of this standard to its implementation status (Met, Partial, Not met, or Deployment responsibility) and a short evidence note.')

# ---------------------------------------------------------------- 5.x NFRs
Cursor(last_req('NFR-DA-PERF-')).req('NFR-DA-PERF-05',
    'With 100 concurrent users on one tenant, the 95th-percentile response time of read endpoints shall stay at or below 1 second, measured by a repeatable load test. A per-user read cache, where used, shall hold an entry for no more than 30 seconds and shall be invalidated by any write in the same Organization.')
c = Cursor(last_req('NFR-DA-SEC-'))
c.req('NFR-DA-SEC-12', 'The login endpoint shall be rate-limited per client, and every response shall carry standard security headers (content-type sniffing protection, frame protection, and a restrictive cross-origin resource policy).')
c.req('NFR-DA-SEC-13', 'Caches, backups, and logs shall never make one tenant\'s data available to another; backups shall be stored outside every web-served directory and be readable only by the service account and the platform administrator.')
Cursor(last_req('NFR-DA-REL-')).req('NFR-DA-REL-05',
    'The backup schedule shall bound data loss to at most 24 hours, and restoring a backup shall be tested at least once per release.')
c = Cursor(last_req('NFR-DA-UX-'))
c.req('NFR-DA-UX-05', 'Every interactive target shall measure at least 24 × 24 CSS pixels (WCAG 2.2 SC 2.5.8), and at least 44 × 44 CSS pixels on touch screens.')
c.req('NFR-DA-UX-06', 'Every release shall pass an automated accessibility scan (e.g. axe-core) of its main screens with no critical or serious violation. Text shall meet WCAG 2.1 AA contrast (4.5:1, or 3:1 for large text); where a brand color falls short, its darker tone shall be used for text, and any remaining brand conflict shall be recorded in Appendix A.')
c.req('NFR-DA-UX-07', 'Every confirmation, warning, and error message shown as a transient notification shall be translated (FR-DA-I18N-07), announced to assistive technologies through a polite live region, and remain readable long enough to act on.')
c = Cursor(last_req('NFR-DA-PORT-'))
c.req('NFR-DA-PORT-03', 'Every change shall be built by a continuous-integration pipeline on Linux, Windows, and macOS that installs dependencies, seeds the demonstration data, runs the automated tests — including the cross-tenant suite of NFR-DA-SEC-11 — and builds the web client.')
c.req('NFR-DA-PORT-04', 'Each tier shall install, seed, and start with its package manager\'s standard commands (install, seed, dev or start), documented in a README at the root of the delivery.')
Cursor(last_req('NFR-DA-MAINT-')).req('NFR-DA-MAINT-06',
    'The Installation Guide and User Guide shall be generated from source in Word and PDF, each with an automatic table of contents, and the User Guide\'s step-by-step walkthroughs shall be replayed automatically against a freshly seeded instance, so the documentation cannot drift from the application.')

# ---------------------------------------------------------------- 6 Data model
add_rows(d.tables[2], [
    ('UserPreference', 'Per-user settings stored server-side: display language, menu dock position, pinned state, collapsed menu groups, favorites, and notification channel choices.'),
    ('BenchmarkMetric / BenchmarkSharing', 'A catalog entry of a comparable metric (unit, direction of good, computation), and an Organization\'s Group-sharing setting, audited on every change.'),
    ('Backup', 'A dated, consistent copy of the data store, with its size and retention expiry; listed to the platform administrator only.'),
    ('RequirementTrace', 'A requirement ID of this standard with its implementation status and evidence note, shown in the traceability view.'),
])

# ---------------------------------------------------------------- 7.1 / 7.2
Cursor(para('Every screen is accessible only when')).body(
    'Every screen sits inside the common navigation shell (Section 4.19). A Benchmarking screen, with Within the '
    'organization and Across the group views, is present in every Dynamic App that holds comparable records (Section 4.20), '
    'and a Backups panel is present on the platform administrator\'s configuration screen (Section 4.21). A screen the user '
    'may not open shows an access notice (FR-DA-RBAC-08) rather than an empty page.')
c = Cursor(para('Alerts: list, run/recompute, mark-read.'))
c.bullet('User preferences: read and update the current user\'s navigation and display preferences.')
c.bullet('Workflow steps: complete, and reopen with justification (FR-DA-AUD-05).')
c.bullet('Benchmarking: metric catalog, within-organization comparison, Group comparison, and read/update of the Organization\'s sharing setting.')
c.bullet('Operations: health (unauthenticated, no tenant data); backup list and on-demand backup (platform administrator).')

# ---------------------------------------------------------------- Appendix A
Cursor(para('This standard states the target platform behavior')).body(
    'Revision 1.1 adds a third reference, the CortexPLM build of September 2026, and rows for the capabilities this '
    'revision introduces. Where a 1.0 target requirement is now observed in CortexPLM, its row says so.')
A = d.tables[3]
seen = ' Revision 1.1: implemented and verified in the CortexPLM reference build.'
for row in ('WBS & Gantt', 'Generic Version Management', 'Multi-Channel Communication', 'Return on Experience', 'External Integration Registry'):
    append_cell(A, row, seen)
append_cell(A, 'Mobile Compatibility', ' Revision 1.1: CortexPLM verified at a 390-pixel viewport with no horizontal scrolling and 44-pixel touch targets; no native wrapper.')
append_cell(A, 'Server-side authorization', ' Revision 1.1: CortexPLM enforces JWT authentication, per-endpoint RBAC, and Pack entitlement server-side, covered by automated tests.')
add_rows(A, [
    ('Navigation Shell & Personalization (Sections 3.11, 4.19)',
     'One shell; menu pinned or sliding, docked on any edge; favorites and collapsed groups stored server-side; RTL mirroring; mobile drawer; keyboard operation.',
     'New in revision 1.1. Implemented and verified in CortexPLM. Neither the journi nor the NCP Solver v2.0 SRS specifies menu personalization; to be assessed.'),
    ('Internal Benchmarking (Sections 3.12, 4.20)',
     'Metrics compared across segments of one Organization and across the Organizations of one Group, aggregates only, minimum sample, opt-out; no external benchmarking.',
     'New in revision 1.1. Implemented in CortexPLM: 13 metrics by project type, track, and department; Group comparison with ranks, median, opt-out, and a minimum of 3 projects. Not present in journi or NCP Solver v2.0.'),
    ('Platform Operations (Sections 3.13, 4.21)',
     'Backups with retention, additive startup migrations, startup guards, health endpoint, deterministic seed, traceability view.',
     'New in revision 1.1. Implemented in CortexPLM (daily backup, 14-day retention). Partial for FR-DA-OPS-05: the runtime version is declared in the package manifest and checked at install time, not by the server at startup. Off-site backup copies remain an operator task (Section 2.7).'),
    ('Opaque identifiers (FR-DA-TEN-11)',
     'Identifiers reveal nothing about other tenants; foreign records answer 404.',
     'New in revision 1.1. Partial in CortexPLM: foreign records answer 404, but record ids are one global sequence, so an id\'s size can hint at overall platform activity.'),
    ('Contrast of brand colors (NFR-DA-UX-06)',
     'WCAG 2.1 AA text contrast; darker brand tone for text where needed.',
     'New in revision 1.1. CortexPLM passes the automated scan except two brand-mandated pairs — white text on the primary orange (2.3:1) and orange key figures (3.0:1) — which conflict with the Brand & Graphical Chart Guide and are recorded here rather than changed.'),
    ('Other 1.1 additions (TEN-09/10, RBAC-08–10, AUD-05, I18N-07/08, VER-06, REX-08, PERF-05, SEC-12/13, REL-05, UX-05/07, PORT-03/04, MAINT-06)',
     'See Appendix B.',
     'New in revision 1.1. All implemented and verified in CortexPLM (load test at 100 users: 95th percentile 0.9 s; CI on Linux, Windows, and macOS).'),
])

# ---------------------------------------------------------------- Appendix B
closing = para('A future revision of either product')
set_runs(closing, ['A future revision of any product\'s SRS that closes one of the gaps above should update this appendix in the same revision, so the standard and the observed state of the suite never drift silently apart.'])
c = Cursor(closing)
c.h1('Appendix B — Revision 1.1 Change Log')
c.body('Revision 1.1 is additive. Every 1.0 requirement keeps its ID and wording, except the editorial changes marked '
       'Clarified below. New requirements take the next free number in their section, so a conformance record written '
       'against 1.0 remains valid against 1.1.')
tbl = copy.deepcopy(A._tbl)
c.el.addnext(tbl)
B = docx.table.Table(tbl, closing._parent)
for tr in B.rows[3:]:
    tbl.remove(tr._tr)
CHANGES = [
    ('Cover, 1.1, 1.5, 1.6', 'Version 1.1; CortexPLM reference', 'Version raised to 1.1. CortexPLM added as a third reference; purpose and overview explain the revision.'),
    ('1.2 Scope', 'Added', 'Navigation shell, internal benchmarking, and platform operations brought into scope; external benchmarking stated out of scope.'),
    ('1.4 Definitions', '6 terms added', 'Navigation Shell, Pinned / Sliding Menu, Dock Position, Internal Benchmarking, Minimum Sample Size, Additive Migration.'),
    ('2.2, 2.3', '3 principles, 3 functions', 'Personal navigation shell; compare within the Group only; upgrade and recover in place.'),
    ('2.5 Operating Environment', 'Clarified', 'Recommended Node.js raised from 22.5+ to 22.13+ (built-in SQLite without an experimental flag); required version declared in the package manifest.'),
    ('2.6, 2.7', '1 constraint, 1 assumption', 'Benchmark data leaving a tenant is aggregated only; off-site backup copies are the operator\'s responsibility.'),
    ('3.1 Tenancy', 'Clarified', 'Group membership is the benchmarking boundary; Group Yes/No and Independent shown on hierarchy screens.'),
    ('3.11 – 3.13', 'New sections', 'Navigation Shell & Personalization, Internal Benchmarking, and Platform Operations architectures.'),
    ('4.1 Tenant & Hierarchy', 'FR-DA-TEN-09 – 11', 'Group membership display; same-tree OBS parents; opaque identifiers and 404 for foreign records.'),
    ('4.2 Identity & RBAC', 'FR-DA-RBAC-08 – 10', 'Access notice instead of a blank page; restricted evaluation-content permission; additive permission migration at startup.'),
    ('4.11 Audit Trail', 'FR-DA-AUD-05', 'Reopen a completed step, with justification, before its stage is decided.'),
    ('4.14 Localization', 'FR-DA-I18N-07 – 08', 'Server messages and assistant answers translated; multilingual Knowledge Base seed with language filter.'),
    ('4.16 Version Management', 'FR-DA-VER-06', 'One shared side-by-side version comparison with changed fields highlighted.'),
    ('4.18 REX', 'FR-DA-REX-08', 'REX prompt shown to whoever closes the record, with Skip that keeps the closure.'),
    ('4.19 Navigation Shell', 'FR-DA-NAV-01 – 10 (new)', 'Menu from one navigation model; collapsible groups; pin / slide away; dock left, right, top, or bottom; favorites; server-side preferences; RTL mirroring; mobile drawer; keyboard operation; standard header bar.'),
    ('4.20 Internal Benchmarking', 'FR-DA-BMK-01 – 09 (new)', 'Metric catalog; within-organization segments; Group ranks with ties; aggregates only; minimum sample; opt-out with audit; permissions; no external benchmarking; charts with filters.'),
    ('4.21 Platform Operations', 'FR-DA-OPS-01 – 08 (new)', 'On-demand and daily backup; retention; restore; additive migrations; startup guards; health endpoint; deterministic seed; traceability view.'),
    ('5.1 Performance', 'NFR-DA-PERF-05', '95th-percentile read time of 1 s or less at 100 concurrent users; rules for a per-user read cache.'),
    ('5.2 Security', 'NFR-DA-SEC-12 – 13', 'Login rate limiting and security headers; no cross-tenant exposure through caches, backups, or logs.'),
    ('5.3 Reliability', 'NFR-DA-REL-05', 'Data loss bounded to 24 hours; restore tested each release.'),
    ('5.5 Usability & Accessibility', 'NFR-DA-UX-05 – 07', 'Target sizes 24 px / 44 px; automated accessibility scan and contrast rule; accessible, translated notifications.'),
    ('5.7 Portability', 'NFR-DA-PORT-03 – 04', 'CI on Linux, Windows, and macOS with seed, tests, and build; standard install / seed / start commands.'),
    ('5.8 Maintainability', 'NFR-DA-MAINT-06', 'Guides generated in Word and PDF with a TOC; walkthroughs replayed against a fresh seed.'),
    ('6 Data Model', '4 entities', 'UserPreference, BenchmarkMetric / BenchmarkSharing, Backup, RequirementTrace.'),
    ('7.1, 7.2 Interfaces', 'Clarified', 'Shell, Benchmarking screen, Backups panel, and access notice; endpoint families for preferences, workflow steps, benchmarking, and operations.'),
    ('Appendix A', 'Updated', 'CortexPLM status added to six 1.0 target rows and server-side authorization; six rows added for 1.1 capabilities; closing note widened to any product.'),
]
hdr = B.rows[0].cells
for cell, text in zip(hdr, ('Section', 'Requirement IDs / Type', 'Change')):
    set_runs(cell.paragraphs[0], [text])
for r, cells in zip(B.rows[1:3], CHANGES[:2]):
    for cell, text in zip(r.cells, cells):
        set_runs(cell.paragraphs[0], [text])
add_rows(B, CHANGES[2:])
WIDTHS = ('2100', '2300', '4500')  # same 8,900 total as the source tables, more room for the change text
for gc, w in zip(tbl.find(qn('w:tblGrid')).findall(qn('w:gridCol')), WIDTHS):
    gc.set(qn('w:w'), w)
for tr in tbl.findall(qn('w:tr')):
    for tc, w in zip(tr.findall(qn('w:tc')), WIDTHS):
        tc.find(qn('w:tcPr')).find(qn('w:tcW')).set(qn('w:w'), w)
c.el = tbl
c.body('Nothing was removed. Product-specific material from CortexPLM — its phase-gate tracks, gate criteria, '
       'industry content, and commercial packaging — was deliberately left out, because it does not apply to every Dynamic App.')

# ---------------------------------------------------------------- manual TOC entries
c = Cursor(para('3.10 Return on Experience', 'Normal'))
for t in ('3.11 Navigation Shell & Personalization Architecture', '3.12 Internal Benchmarking Architecture', '3.13 Platform Operations Architecture'):
    c.add(T_TOC2, [t, '0'])
c = Cursor(para('4.18 Return on Experience', 'Normal'))
for t in ('4.19 Navigation Shell & Personalization', '4.20 Internal Benchmarking', '4.21 Platform Operations'):
    c.add(T_TOC2, [t, '0'])
Cursor(T_TOC1).add(T_TOC1, ['Appendix B — Revision 1.1 Change Log', '0'])

d.core_properties.revision = (d.core_properties.revision or 1) + 1
d.save(OUT)
print('saved', OUT)
