import copy, re
from docx.oxml.ns import qn
d = None
def paras():
    return d.paragraphs

def find(prefix, style=None):
    for p in d.paragraphs:
        if p.text.startswith(prefix) and (style is None or p.style.name == style):
            return p
    raise KeyError(prefix)

def strip_marks(el):
    for tag in ('w:bookmarkStart', 'w:bookmarkEnd', 'w:proofErr'):
        for b in el.iter(qn(tag)):
            pass
        for b in list(el.iter(qn(tag))):
            b.getparent().remove(b)
    return el

def set_runs(p_el, texts):
    """Keep the first len(texts) runs (formatting preserved), set their text, drop the rest."""
    runs = [r for r in p_el.iter(qn('w:r'))]
    for i, r in enumerate(runs):
        if i < len(texts):
            ts = r.findall(qn('w:t'))
            for t in ts[1:]:
                r.remove(t)
            for tab in r.findall(qn('w:tab')):
                r.remove(tab)
            if not ts:
                t = r.makeelement(qn('w:t'), {}); r.append(t); ts = [t]
            ts[0].text = texts[i]
            ts[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        else:
            r.getparent().remove(r)
    return p_el

def clone_para(model, texts):
    el = strip_marks(copy.deepcopy(model._p if hasattr(model, '_p') else model))
    return set_runs(el, texts)

def insert_after(ref, els):
    ref = ref._p if hasattr(ref, '_p') else ref
    for el in els:
        ref.addnext(el); ref = el
    return ref



def init():
    global M_NORMAL, M_LIST, M_REQ, M_H2, M_H1
    M_NORMAL = find('Every Dynamic App offers a catalog of project templates', 'Normal')
    M_LIST = find('Upgrade and recover in place', 'List Paragraph')
    M_REQ = find('FR-DA-QLT-12', 'List Paragraph')
    M_H2 = find('4.36 Questionnaire', 'Heading 2')
    M_H1 = find('Appendix F', 'Heading 1')
def req(i, t): return clone_para(M_REQ, [i + ': ', t])
def normal(t): return clone_para(M_NORMAL, [t])
def bullet(t): return clone_para(M_LIST, [t])
def h2(t): return clone_para(M_H2, [t])
def table_with(header0):
    for t in d.tables:
        if t.rows[0].cells[0].text.strip() == header0:
            yield t
def set_cell(tc, text):
    ps = tc.findall(qn('w:p'))
    for p in ps[1:]: tc.remove(p)
    p = ps[0]
    runs = list(p.iter(qn('w:r')))
    if not runs:
        r = p.makeelement(qn('w:r'), {}); p.append(r)
    set_runs(p, [text])
def add_rows(t, rows):
    tbl = t._tbl
    for vals in rows:
        trs = tbl.findall(qn('w:tr'))
        model = trs[-2] if len(trs) > 2 else trs[-1]
        tr = strip_marks(copy.deepcopy(model))
        for tc, v in zip(tr.findall(qn('w:tc')), vals):
            set_cell(tc, v)
        trs[-1].addnext(tr)
def add_tab_run(el):
    runs = list(el.iter(qn('w:r')))
    r = runs[1]
    for t in r.findall(qn('w:t')): r.remove(t)
    tab = r.makeelement(qn('w:tab'), {}); r.append(tab)
    t = r.makeelement(qn('w:t'), {}); t.text = '0'; r.append(t)
    return el
def toc(model_prefix, text):
    return add_tab_run(clone_para(find(model_prefix, 'Normal'), [text, '0']))
