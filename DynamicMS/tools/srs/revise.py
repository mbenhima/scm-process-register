import copy, sys, re
import docx
from docx.oxml.ns import qn
sys.path.insert(0, __import__('os').path.dirname(__import__('os').path.abspath(__file__)))
from content15 import *

SRC, OUT = sys.argv[1], sys.argv[2]
d = docx.Document(SRC)
body = d.element.body

def paras():
    return d.paragraphs

def find(prefix, style=None):
    for p in d.paragraphs:
        if p.text.startswith(prefix) and (style is None or p.style.name == style):
            return p
    raise KeyError(prefix)

def strip_marks(el):
    for tag in ('w:bookmarkStart', 'w:bookmarkEnd', 'w:proofErr'):
        for b in el.iter(qn(tag)):
            pass
        for b in list(el.iter(qn(tag))):
            b.getparent().remove(b)
    return el

def set_runs(p_el, texts):
    """Keep the first len(texts) runs (formatting preserved), set their text, drop the rest."""
    runs = [r for r in p_el.iter(qn('w:r'))]
    for i, r in enumerate(runs):
        if i < len(texts):
            ts = r.findall(qn('w:t'))
            for t in ts[1:]:
                r.remove(t)
            for tab in r.findall(qn('w:tab')):
                r.remove(tab)
            if not ts:
                t = r.makeelement(qn('w:t'), {}); r.append(t); ts = [t]
            ts[0].text = texts[i]
            ts[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        else:
            r.getparent().remove(r)
    return p_el

def clone_para(model, texts):
    el = strip_marks(copy.deepcopy(model._p if hasattr(model, '_p') else model))
    return set_runs(el, texts)

def insert_after(ref, els):
    ref = ref._p if hasattr(ref, '_p') else ref
    for el in els:
        ref.addnext(el); ref = el
    return ref

# models
M_NORMAL = find('Every Dynamic App offers a catalog of project templates', 'Normal')
M_LIST = find('Upgrade and recover in place', 'List Paragraph')           # plain bullet
M_REQ = find('FR-DA-QLT-12', 'List Paragraph')                             # bold id + text
M_H2 = find('4.36 Questionnaire', 'Heading 2')
M_H1 = find('Appendix F', 'Heading 1')
M_REF = find('CortexSkills Specs Pack', 'List Paragraph')

def req(i, t): return clone_para(M_REQ, [i + ': ', t])
def normal(t): return clone_para(M_NORMAL, [t])
def bullet(t): return clone_para(M_LIST, [t])
def h2(t): return clone_para(M_H2, [t])

# --- cover and introduction
cov = find('Version 1.4', 'Normal'); set_runs(cov._p, [cov.text.replace('1.4', '1.5')])
insert_after(find('Revision 1.4 adds questionnaire'), [normal(PURPOSE)])
insert_after(find('Revision 1.4 adds to this scope'), [normal(SCOPE)])
insert_after(M_REF, [clone_para(M_REF, [REFERENCE])])
insert_after(find('Revision 1.4 adds Section 4.36'), [normal(OVERVIEW)])
insert_after(M_LIST, [bullet(t) for t in PRINCIPLES])
insert_after(find('Design questionnaires from a versioned', 'List Paragraph'), [bullet(t) for t in FUNCTIONS])

# --- architecture 3.15 - 3.18 after the last paragraph of 3.14
anchor = find('Activation is additive and ordered')
els = []
for title, ps in ARCH:
    els.append(h2(title)); els += [normal(t) for t in ps]
insert_after(anchor, els)

# --- section 4 intro
insert_after(find('Requirements added in revision 1.3 in Sections 4.25'),
             [normal('Requirements added in revision 1.5 are generalized from the DynamicMS reference build and from its user '
                     'feedback; they carry new IDs, never reuse an earlier one, and are listed in Appendix G.')])

# --- appended requirements
for anchor_id, items in APPEND:
    insert_after(find(anchor_id + ':'), [req(i, t) for i, t in items])

# --- new sections 4.37 - 4.43
els = []
for title, intro, items in SECTIONS:
    els.append(h2(title)); els.append(normal(intro)); els += [req(i, t) for i, t in items]
insert_after(find('FR-DA-QLT-12:'), els)

# --- tables
def table_with(header0):
    for t in d.tables:
        if t.rows[0].cells[0].text.strip() == header0:
            yield t

def set_cell(tc, text):
    ps = tc.findall(qn('w:p'))
    for p in ps[1:]: tc.remove(p)
    p = ps[0]
    runs = list(p.iter(qn('w:r')))
    if not runs:
        r = p.makeelement(qn('w:r'), {}); p.append(r); runs = [r]
    set_runs(p, [text])

def add_rows(t, rows):
    tbl = t._tbl
    for vals in rows:
        trs = tbl.findall(qn('w:tr'))
        model = trs[-2] if len(trs) > 2 else trs[-1]
        tr = strip_marks(copy.deepcopy(model))
        for tc, v in zip(tr.findall(qn('w:tc')), vals):
            set_cell(tc, v)
        trs[-1].addnext(tr)

terms = next(t for t in table_with('Term'))
add_rows(terms, TERMS)
ent = next(t for t in table_with('Entity'))
add_rows(ent, ENTITIES)
area = next(t for t in table_with('Area'))
add_rows(area, CONFORMANCE)

# --- 7.1 paragraph
insert_after(find('Revision 1.3 adds screens for verticals'), [normal(UI_PARA)])

# --- Appendix G after the Appendix F table
ftab = [t for t in table_with('Section')][-1]
gtab = strip_marks(copy.deepcopy(ftab._tbl))
trs = gtab.findall(qn('w:tr'))
for tr in trs[1:]: gtab.remove(tr)
g_h = clone_para(M_H1, ['Appendix G — Revision 1.5 Change Log'])
g_i = normal('Revision 1.5 keeps every earlier requirement ID and is purely additive. Its requirements are application '
             'agnostic: they come from the DynamicMS reference build and from the product direction received with its '
             'user feedback, restated so that any Dynamic App can meet them.')
insert_after(ftab._tbl, [g_h, g_i, gtab])
from docx.table import Table
gt = Table(gtab, d._body)
for vals in CHANGELOG:
    trs = gtab.findall(qn('w:tr'))
    model = ftab._tbl.findall(qn('w:tr'))[1 + (len(trs) - 1) % 2]
    tr = strip_marks(copy.deepcopy(model))
    for tc, v in zip(tr.findall(qn('w:tc')), vals): set_cell(tc, v)
    trs[-1].addnext(tr)

# --- table of contents (static entries; page numbers fixed afterwards)
def toc_entry(model_prefix, text):
    m = find(model_prefix, 'Normal')
    return clone_para(m, [text, '\t0'])

def add_tab_run(el):
    # second run holds the tab + page number
    runs = list(el.iter(qn('w:r')))
    r = runs[1]
    for t in r.findall(qn('w:t')): r.remove(t)
    tab = r.makeelement(qn('w:tab'), {}); r.append(tab)
    t = r.makeelement(qn('w:t'), {}); t.text = '0'; r.append(t)
    return el

def toc(model_prefix, text):
    return add_tab_run(clone_para(find(model_prefix, 'Normal'), [text, '0']))

insert_after(find('3.14 Vertical & SME Layer Architecture\t', 'Normal'), [toc('3.14 Vertical', t) for t, _ in ARCH])
insert_after(find('4.36 Questionnaire & Survey Management\t', 'Normal'), [toc('4.36 Questionnaire', s[0]) for s in SECTIONS])
insert_after(find('Appendix F — Revision 1.4 Change Log\t', 'Normal'), [toc('Appendix F', 'Appendix G — Revision 1.5 Change Log')])

d.save(OUT)
print('saved', OUT)
