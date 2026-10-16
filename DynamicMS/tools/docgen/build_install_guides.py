"""Builds the French installation guides of the DynamicMS runtime (Word with TOC; PDF by lo_convert):
  customer — Guide_Installation_Docker_Runtime_Customer_Linux_DynamicMS_FR.docx
  editor   — Guide_Installation_Docker_Runtime_Editor_Linux_Windows_DynamicMS_FR.docx
Screenshots: deliverables/build/runtime-fr (taken on the runtime images).
Usage: python3 build_install_guides.py [customer|editor|all]"""
import json
import os
import sys

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

import brand
from brand import new_document, cover, toc, page_break, para, bullets, numbered, table, callout, image, rgb, NAVY, INK, MUTED, BG

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SHOTS = os.path.join(ROOT, 'deliverables', 'build', 'runtime-fr')
OUT = os.path.join(ROOT, 'deliverables')
VERSION = '2026.10'
FP = 'da55:898d:40dd:9909'
CUST_IMG = f'Docker_Runtime_Customer_Linux_DynamicMS_{VERSION}_Image.tar.gz'
ED_IMG = f'Docker_Runtime_Editor_Linux_DynamicMS_{VERSION}_Image.tar.gz'
CUST_ZIP = f'Docker_Runtime_Customer_Linux_DynamicMS_{VERSION}_Deploy.zip'
ED_LIN_ZIP = f'Docker_Runtime_Editor_Linux_DynamicMS_{VERSION}_Deploy.zip'
ED_WIN_ZIP = f'Docker_Runtime_Editor_Windows_DynamicMS_{VERSION}_Deploy.zip'
KEY = 'Docker_Runtime_Editor_Licence_PrivateKey.pem'
EVAL = f'Licence_Evaluation_Customer_DynamicMS_{VERSION}.lic'
G_CUST = 'Guide_Installation_Docker_Runtime_Customer_Linux_DynamicMS_FR'
G_ED = 'Guide_Installation_Docker_Runtime_Editor_Linux_Windows_DynamicMS_FR'
G_USER = 'Guide_Utilisateur_Runtime_Customer_DynamicMS_FR'


def code(doc, lines):
    """Commands in a monospace block on a light background."""
    for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(4 if i == 0 else 0); pf.space_after = Pt(4 if i == len(lines) - 1 else 0)
        pf.left_indent = Cm(0.3); pf.line_spacing = 1.0
        ppr = p._p.get_or_add_pPr()
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), 'EEF2F6'); ppr.append(shd)
        r = p.add_run(line if line else ' ')
        r.font.name = 'Consolas'; r.font.size = Pt(8.5); r.font.color.rgb = rgb(NAVY)
        rpr = r._r.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts'); rpr.append(rf)
        for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
            rf.set(qn(a), 'Consolas')


def shot(doc, name, cap, width=16):
    image(doc, os.path.join(SHOTS, f'{name}.png'), width, cap)


def h1(doc, t):
    doc.add_heading(t, level=1)


def h2(doc, t):
    doc.add_heading(t, level=2)


def h3(doc, t):
    doc.add_heading(t, level=3)


def docker_install(doc):
    h3(doc, 'Ubuntu 22.04 / 24.04 LTS et Debian 12')
    code(doc, [
        'sudo apt-get update && sudo apt-get install -y ca-certificates curl',
        'sudo install -m 0755 -d /etc/apt/keyrings',
        'sudo curl -fsSL https://download.docker.com/linux/ubuntu/gpg -o /etc/apt/keyrings/docker.asc',
        'echo "deb [arch=$(dpkg --print-architecture) signed-by=/etc/apt/keyrings/docker.asc] \\',
        '  https://download.docker.com/linux/ubuntu $(. /etc/os-release && echo $VERSION_CODENAME) stable" \\',
        '  | sudo tee /etc/apt/sources.list.d/docker.list',
        'sudo apt-get update',
        'sudo apt-get install -y docker-ce docker-ce-cli containerd.io docker-compose-plugin',
        'sudo systemctl enable --now docker',
    ])
    para(doc, 'Sous Debian, remplacer « ubuntu » par « debian » dans les deux adresses.', size=9.5, italic=True)
    h3(doc, 'Red Hat Enterprise Linux 9, Rocky Linux 9, AlmaLinux 9')
    code(doc, [
        'sudo dnf -y install dnf-plugins-core',
        'sudo dnf config-manager --add-repo https://download.docker.com/linux/rhel/docker-ce.repo',
        'sudo dnf -y install docker-ce docker-ce-cli containerd.io docker-compose-plugin',
        'sudo systemctl enable --now docker',
    ])
    h3(doc, 'Contrôle')
    code(doc, ['docker --version            # 24 ou supérieur', 'docker compose version      # v2.x', 'sudo docker run --rm hello-world   # facultatif, nécessite Internet'])
    para(doc, 'Serveur sans accès Internet : installer les paquets Docker téléchargés au préalable (dépôt miroir interne ou paquets .deb/.rpm copiés sur le serveur). DynamicMS lui-même ne nécessite aucun accès Internet : l\'image contient tout ce qu\'il faut.', size=10)


def env_table(doc, editor=False):
    rows = [
        ['ADMIN_EMAIL', 'E-mail de l\'administrateur de la plateforme, créé au premier démarrage.', 'Oui'],
        ['ADMIN_NAME', 'Nom affiché de cet administrateur.', 'Non'],
        ['ADMIN_PASSWORD', 'Mot de passe provisoire (12 caractères, majuscule, minuscule, chiffre), à changer à la première connexion. Utilisé une seule fois.', 'Oui'],
        ['SEED_PASSWORD', 'Mot de passe initial de tous les comptes des données initialisées ; chacun le change à sa première connexion.', 'Oui (INIT_DATA=demo)'],
        ['JWT_SECRET', 'Secret de signature des sessions, 32 caractères au moins (généré par l\'installation).', 'Oui'],
        ['INIT_DATA', 'demo : données initialisées de tous les secteurs et des PME ; empty : catalogue et administrateur seulement. Lu au premier démarrage uniquement.', 'Non (demo)'],
        ['CORS_ORIGIN', 'Adresse publique de l\'application (https://…).', 'Oui'],
        ['HTTP_BIND / HTTP_PORT', 'Adresse et port d\'écoute sur le serveur (127.0.0.1:4000 par défaut, derrière le proxy HTTPS).', 'Non'],
        ['BACKUP_RETENTION_DAYS', 'Durée de conservation des sauvegardes quotidiennes internes (jours).', 'Non'],
    ]
    if editor:
        rows.append(['LICENCE_SIGNING_KEY', 'Chemin de la clé privée dans le conteneur (fixé par docker-compose.yml : /run/secrets/licence_private.pem).', 'Fixé'])
    table(doc, ['Variable', 'Rôle', 'Obligatoire'], rows, [4.0, 10.0, 3.0], size=9)


def seeded_accounts(doc):
    orgs = json.load(open(os.path.join(SHOTS, 'orgs.json'), encoding='utf8'))
    segs = {s['id']: s['fr'] for s in json.load(open(os.path.join(SHOTS, 'segs.json'), encoding='utf8'))}
    roles = json.load(open(os.path.join(SHOTS, 'roles.json'), encoding='utf8'))
    slugs = [('ims', 'ims_manager'), ('quality', 'quality_manager'), ('hse', 'hse_manager'), ('risk', 'risk_manager'), ('compliance', 'compliance_officer'), ('audit', 'audit_manager'), ('hr', 'hr_manager'), ('documents', 'document_controller'), ('it', 'it_manager'), ('operations', 'operations_manager'), ('performance', 'performance_manager'), ('esg', 'esg_manager'), ('transformation', 'transformation_manager'), ('process', 'process_excellence_manager'), ('admin', 'tenant_admin'), ('ceo', 'top_management'), ('employee', 'employee'), ('aigov', 'ai_governance_officer'), ('auditor', 'auditor')]
    para(doc, f'Les données initialisées couvrent les 29 secteurs et le profil universel, chacun avec une grande entreprise et une PME : {len(orgs)} organisations, 120 projets complets (SMQ et QHSE), 1 140 comptes. Chaque organisation possède un compte par rôle, de la forme rôle@domaine. Tous les comptes se connectent la première fois avec SEED_PASSWORD et choisissent alors leur propre mot de passe.')
    h3(doc, 'Comptes de chaque organisation')
    table(doc, ['Préfixe de l\'e-mail', 'Rôle'], [[f'{s}@<domaine>', roles.get(r, r)] for s, r in slugs], [5, 12], size=9)
    h3(doc, 'Organisations et domaines')
    table(doc, ['Secteur', 'Organisation', 'Taille', 'Domaine des comptes'], [[f"{o['sector']} — {segs.get(o['sector'], 'Universel') if o['sector'] != 'UNI' else 'Universel'}", o['fr'], 'Grande' if o['size'] == 'Large' else 'PME', o['email_domain']] for o in orgs], [4.6, 5.4, 1.6, 5.4], size=8)
    para(doc, 'Exemple : quality@horizon-universal.example est le responsable qualité de Horizon Holding Universelle ; ims@atlas-sme.example est le responsable SMI de la PME Atlas.', size=9.5, italic=True)


def operations(doc, svc, vol, editor=False, windows=False):
    h2(doc, 'Démarrer, arrêter, consulter les journaux')
    code(doc, ['cd dynamicms-' + ('editor' if editor else 'customer'), 'sudo docker compose up -d        # démarrer', 'sudo docker compose stop         # arrêter (données conservées)', f'sudo docker compose ps           # état ; « healthy » quand l\'application répond', f'sudo docker logs -f {svc}   # journaux', 'curl -s http://127.0.0.1:4000/api/health   # santé'])
    h2(doc, 'Sauvegardes')
    bullets(doc, [
        ('Sauvegardes internes : ', f'l\'application copie chaque jour sa base dans le volume {vol} (dossier backups) et les conserve BACKUP_RETENTION_DAYS jours. Administration › Exploitation › Sauvegarder maintenant en crée une à la demande.'),
        ('Sauvegarde complète : ', 'backup.sh arrête l\'application quelques secondes et archive tout le volume (base, pièces jointes, licence, sauvegardes internes) dans ./sauvegardes. À planifier chaque nuit et à copier hors du serveur.'),
        ('Restauration : ', 'restore.sh remplace toutes les données par une archive de backup.sh (confirmation RESTAURER demandée).'),
    ], size=10)
    code(doc, ['sudo ./backup.sh', 'sudo ./restore.sh sauvegardes/' + ('dynamicms-editor-data' if editor else 'dynamicms-data') + '-AAAAMMJJ-HHMMSS.tar.gz', '# planification nocturne (crontab -e, en root) :', f'30 2 * * * cd /opt/dynamicms/dynamicms-{"editor" if editor else "customer"} && ./backup.sh >> /var/log/dynamicms-backup.log 2>&1'])
    h2(doc, 'Mise à jour vers une nouvelle version')
    numbered(doc, [
        'Faire une sauvegarde complète (backup.sh).',
        'Copier la nouvelle archive d\'image et son fichier .sha256, vérifier l\'empreinte, puis : sudo docker load -i <archive>.',
        'Remplacer la version de l\'image dans docker-compose.yml (ligne image:) si elle a changé, puis : sudo docker compose up -d.',
        'Au démarrage, les données sont migrées sur place (aucune réinitialisation) ; vérifier les journaux et la santé.',
    ])


def troubleshooting(doc, editor=False):
    rows = [
        ['Refusing to start: set ADMIN_EMAIL and ADMIN_PASSWORD…', 'Premier démarrage sans administrateur.', 'Compléter .env puis docker compose up -d.'],
        ['…needs at least 12 characters…', 'Mot de passe trop faible (ADMIN_PASSWORD ou SEED_PASSWORD).', 'Choisir 12 caractères avec majuscule, minuscule et chiffre.'],
        ['Refusing to start: set JWT_SECRET…', 'Secret absent ou trop court.', 'openssl rand -hex 32, à reporter dans .env.'],
        ['Page de connexion inaccessible', 'Proxy HTTPS ou pare-feu ; application encore en initialisation.', 'curl http://127.0.0.1:4000/api/health sur le serveur ; docker logs.'],
        ['« Changez votre mot de passe » à chaque requête', 'Compte avec mot de passe initial.', 'Normal : choisir son mot de passe sur l\'écran proposé.'],
        ['Archive d\'image : empreinte incorrecte', 'Fichier incomplet ou modifié pendant le transfert.', 'Copier à nouveau l\'archive (transfert binaire) et revérifier.'],
        ['Port 4000 déjà utilisé', 'Autre service sur le serveur.', 'Changer HTTP_PORT dans .env et le proxy.'],
    ]
    if not editor:
        rows += [['Plateforme en lecture seule', 'Aucune licence valide (absente, expirée, sièges dépassés).', 'Administration › Exploitation › Licence ; demander un fichier à l\'éditeur.'],
                 ['La signature de la licence n\'est pas valide', 'Fichier modifié, ou émis pour une autre installation de l\'éditeur.', 'Ne jamais modifier le fichier ; demander un nouvel envoi à l\'éditeur.']]
    else:
        rows += [['Clé de signature : « The key cannot be read: EACCES »', 'Droits du fichier de clé.', 'sudo chown 1000:1000 secrets/licence_private.pem && sudo chmod 400 secrets/licence_private.pem'],
                 ['« Cette clé ne correspond pas… »', 'Clé différente de celle intégrée aux images clientes.', 'Utiliser la clé fournie (empreinte ' + FP + ').']]
    table(doc, ['Symptôme', 'Cause', 'Action'], rows, [5.2, 5.4, 6.4], size=8.5)


# ======================================================================= customer guide
def customer():
    brand.set_lang('fr')
    doc = new_document('Guide d\'installation — Runtime client Linux', 'DynamicMS — Guide d\'installation du runtime client (Linux)')
    cover(doc, 'Guide d\'installation', 'Runtime client Linux', 'DynamicMS 2026.10 — cloud privé ou serveur sur site, avec activation de la licence',
          'AI Value · DynamicMS 2026.10 · octobre 2026')
    toc(doc)

    h1(doc, 'À propos de ce guide')
    para(doc, 'Ce guide s\'adresse à l\'administrateur système qui installe DynamicMS sur l\'infrastructure du client : un serveur Linux dans le centre de données du client ou une machine virtuelle de son cloud privé. Il couvre la préparation du serveur, l\'installation du runtime, le premier accès, l\'activation de la licence, l\'exploitation courante et le dépannage.')
    bullets(doc, [('Runtime : ', 'l\'application est livrée exécutable, sans code source. Version de production optimisée : le serveur est minifié puis compilé en code machine virtuel (bytecode V8) ; l\'interface web est minifiée, découpée en modules chargés à la demande et précompressée (Brotli, gzip).'),
                  ('Données initialisées : ', 'au premier démarrage, l\'application charge les données de tous les secteurs et des PME (chapitre 8), utiles à l\'évaluation et à la formation.'),
                  ('Mode production : ', 'aucun compte n\'est affiché sur la page de connexion ; chaque compte reçoit un mot de passe initial qu\'il remplace à sa première connexion.'),
                  ('Conventions : ', 'les commandes sont à exécuter sur le serveur, dans un terminal, avec un compte disposant de sudo.')], size=10)

    h1(doc, 'Contenu de la livraison')
    para(doc, 'Tous les fichiers cités dans ce guide sont fournis ci-dessous. Les copier dans un même dossier du serveur, par exemple /opt/dynamicms.')
    table(doc, ['Fichier', 'Contenu'], [
        [CUST_IMG, 'Image du runtime client (conteneur Linux x86_64).'],
        [CUST_IMG + '.sha256', 'Empreinte SHA-256 de l\'archive, pour vérifier son intégrité.'],
        [CUST_ZIP, 'Dossier dynamicms-customer : docker-compose.yml, .env.example, install.sh, backup.sh, restore.sh, nginx/dynamicms.conf, licence/' + EVAL + ', LISEZMOI.txt.'],
        [G_CUST + '.docx / .pdf', 'Le présent guide.'],
        [G_USER + '.docx / .pdf', 'Guide utilisateur : déroulement complet des processus de bout en bout.'],
        ['Licence_LIC-AAAA-NNN_vN_<Client>.lic', 'Licence définitive, émise par l\'éditeur pour le client (chapitre 7).'],
    ], [7.2, 9.8], size=9)
    h2(doc, 'Contenu du dossier dynamicms-customer')
    table(doc, ['Fichier', 'Rôle'], [
        ['docker-compose.yml', 'Définition du service : image, volume de données dynamicms-data, port, journaux.'],
        ['.env.example', 'Modèle des paramètres ; install.sh le copie en .env (droits 600).'],
        ['install.sh', 'Installation : contrôle de Docker, vérification de l\'empreinte, chargement de l\'image, création de .env, démarrage, attente de l\'application.'],
        ['backup.sh / restore.sh', 'Sauvegarde complète du volume de données et restauration.'],
        ['nginx/dynamicms.conf', 'Exemple de proxy inverse HTTPS (nginx).'],
        ['licence/' + EVAL, 'Licence d\'évaluation valable jusqu\'au 31 janvier 2027, 1 500 comptes actifs.'],
        ['LISEZMOI.txt', 'Résumé du contenu et de l\'installation.'],
    ], [5.5, 11.5], size=9)

    h1(doc, 'Ce que contient le runtime')
    table(doc, ['Élément', 'Description'], [
        ['Application', 'Serveur DynamicMS minifié et compilé (bytecode V8), interface web minifiée et précompressée ; aucun fichier source.'],
        ['Mode', 'Production, sur site : la licence est exigée pour toute modification des données.'],
        ['Données', 'Créées au premier démarrage dans le volume dynamicms-data : catalogue (12 processus de bout en bout, 162 macro-processus), données initialisées de tous les secteurs et des PME.'],
        ['Licence', 'Vérifiée par la clé publique de l\'éditeur intégrée à l\'application ; aucune autre clé n\'est acceptée.'],
        ['Langues', 'Français, anglais et arabe (de droite à gauche).'],
        ['Accès réseau', 'Aucun accès Internet requis. Les modèles d\'IA externes restent désactivés tant qu\'un administrateur n\'en configure pas.'],
    ], [3.5, 13.5], size=9.5)

    h1(doc, 'Prérequis')
    table(doc, ['Élément', 'Minimum', 'Recommandé'], [
        ['Système', 'Linux x86_64 : Ubuntu 22.04/24.04 LTS, Debian 12, RHEL/Rocky/AlmaLinux 9', 'Ubuntu 24.04 LTS'],
        ['Processeur', '2 vCPU', '4 vCPU'],
        ['Mémoire', '4 Go', '8 Go'],
        ['Disque', '20 Go libres', '50 Go (pièces jointes, sauvegardes)'],
        ['Logiciels', 'Docker Engine 24, Docker Compose v2, openssl, curl', 'Versions à jour'],
        ['Réseau', 'Port 443 ouvert vers le proxy HTTPS ; port 4000 local', 'Nom DNS et certificat TLS de l\'entreprise'],
        ['Heure', 'Horloge synchronisée (NTP) : la date d\'expiration de la licence en dépend', 'chrony ou systemd-timesyncd'],
    ], [3.0, 8.0, 6.0], size=9)
    h2(doc, 'Installer Docker')
    docker_install(doc)

    h1(doc, 'Installation pas à pas')
    h2(doc, 'Copier et vérifier les fichiers')
    code(doc, ['sudo mkdir -p /opt/dynamicms && cd /opt/dynamicms', '# copier ici les fichiers livrés (scp, clé USB…), puis :', f'sha256sum -c {CUST_IMG}.sha256', f'# résultat attendu : {CUST_IMG}: OK', f'sudo apt-get install -y unzip   # ou : sudo dnf install -y unzip', f'unzip {CUST_ZIP}'])
    h2(doc, 'Lancer l\'installation')
    code(doc, ['cd /opt/dynamicms/dynamicms-customer', 'sudo ./install.sh'])
    para(doc, 'install.sh enchaîne les opérations suivantes et s\'arrête avec un message explicite en cas de problème :')
    numbered(doc, [
        'contrôle de Docker et de Docker Compose ;',
        'vérification de l\'empreinte SHA-256 de l\'archive, puis chargement de l\'image dynamicms-customer:2026.10 ;',
        'création de .env (droits 600) avec un JWT_SECRET aléatoire ;',
        'saisie de l\'e-mail et du mot de passe provisoire de l\'administrateur, puis du mot de passe initial des comptes initialisés (saisie masquée, règle des 12 caractères contrôlée) ;',
        'démarrage du service et attente de l\'application : le premier démarrage crée les données (environ une minute).',
    ])
    h2(doc, 'Paramètres du fichier .env')
    para(doc, 'Les valeurs peuvent aussi être renseignées dans .env avant de lancer install.sh, qui ne pose alors plus de question.')
    env_table(doc)
    callout(doc, 'ADMIN_PASSWORD, SEED_PASSWORD et INIT_DATA ne sont lus qu\'au premier démarrage, quand le volume de données est vide. Les modifier ensuite n\'a pas d\'effet. Une fois les mots de passe changés à la première connexion, ces deux valeurs peuvent être effacées de .env.', 'Bon à savoir —')
    h2(doc, 'Vérifier l\'installation')
    code(doc, ['sudo docker compose ps              # STATUS : Up … (healthy)', 'curl -s http://127.0.0.1:4000/api/health', '# {"status":"ok",…,"mode":"production","demo":false,…,"deployment":"onprem","licence":"invalid"}'])
    para(doc, '« licence »: « invalid » est normal à ce stade : la licence s\'installe au chapitre 7.', size=10, italic=True)
    h2(doc, 'Publier en HTTPS')
    para(doc, 'Par défaut, l\'application n\'écoute que sur le serveur lui-même (127.0.0.1:4000). Le proxy inverse de l\'entreprise la publie en HTTPS. Exemple avec nginx (fichier fourni nginx/dynamicms.conf) :')
    code(doc, ['sudo apt-get install -y nginx            # ou : sudo dnf install -y nginx', 'sudo cp nginx/dynamicms.conf /etc/nginx/conf.d/dynamicms.conf', '# adapter server_name, ssl_certificate et ssl_certificate_key', 'sudo nginx -t && sudo systemctl reload nginx', '# pare-feu : sudo ufw allow 443/tcp   ou   sudo firewall-cmd --add-service=https --permanent && sudo firewall-cmd --reload'])
    para(doc, 'Renseigner ensuite CORS_ORIGIN avec l\'adresse publique (https://dynamicms.client.local) et appliquer : sudo docker compose up -d.', size=10)

    h1(doc, 'Premier accès')
    para(doc, 'Ouvrir l\'adresse publique dans un navigateur (Chrome, Edge, Firefox ou Safari récents). La page de connexion ne montre aucun compte : chaque utilisateur doit connaître son e-mail et son mot de passe.')
    shot(doc, 'c-login', 'Figure — Page de connexion du runtime : aucun compte de démonstration n\'est affiché.', 15)
    para(doc, 'Se connecter avec ADMIN_EMAIL et le mot de passe provisoire. L\'application demande aussitôt un nouveau mot de passe ; aucune autre action n\'est possible avant.')
    shot(doc, 'c-firstpw', 'Figure — Choix du mot de passe à la première connexion.', 15)

    h1(doc, 'Activation de la licence')
    h2(doc, 'Principe')
    bullets(doc, [
        ('Fichier de licence : ', 'un fichier .lic émis par l\'éditeur pour le client. Il indique le client, le nombre de comptes actifs autorisés (postes), la date d\'expiration et les éléments sous licence, et il est signé par l\'éditeur.'),
        ('Vérification : ', 'l\'application vérifie la signature avec la clé publique de l\'éditeur intégrée au runtime. Un fichier modifié, même d\'un caractère, est refusé.'),
        ('Sans licence valide : ', 'la plateforme est en lecture seule pour tous les comptes, administrateur compris ; seule l\'installation d\'une licence reste possible.'),
        ('Avec licence valide : ', 'seuls les éléments sous licence peuvent être activés pour une organisation, et le nombre de comptes actifs ne peut pas dépasser les postes.'),
    ], size=10)
    h2(doc, 'Installer la licence')
    numbered(doc, [
        'Se connecter avec le compte administrateur de la plateforme.',
        'Ouvrir Administration › onglet Exploitation. La carte Licence indique « ÉCHEC — Aucune licence installée ».',
        'Sélectionner Installer un fichier de licence et choisir le fichier .lic (licence d\'évaluation : dossier licence/ du paquet ; licence définitive : fichier reçu de l\'éditeur).',
        'La carte passe à « VALIDE » et affiche le client, les postes, la date d\'expiration et les éléments sous licence. La plateforme est utilisable.',
    ])
    shot(doc, 'c-licence-before', 'Figure — Avant l\'activation : aucune licence installée, plateforme en lecture seule.', 15)
    shot(doc, 'c-licence-after', 'Figure — Après l\'activation : licence valide, avec ses postes, son expiration et ses éléments.', 15)
    h2(doc, 'Obtenir la licence définitive')
    para(doc, 'Demander à l\'éditeur le fichier de licence définitif en indiquant : le nom du client, un contact technique, le nombre de comptes actifs prévus (le jeu de données initialisé en compte 1 140), la date de fin souhaitée et le mode de déploiement (cloud privé ou sur site). L\'éditeur renvoie un fichier nommé Licence_LIC-AAAA-NNN_vN_<Client>.lic.')
    h2(doc, 'Renouveler ou modifier la licence')
    para(doc, 'À chaque changement (postes, date, éléments), l\'éditeur émet une nouvelle version du fichier. Il suffit de l\'installer de la même façon : elle remplace la précédente. Si le nouveau fichier n\'est pas valide, il est refusé et la licence en place est conservée.')
    h2(doc, 'Messages de licence')
    table(doc, ['Message', 'Signification', 'Action'], [
        ['Licence valide.', 'La licence est installée et vérifiée.', '—'],
        ['Aucune licence installée.', 'Premier démarrage ou fichier retiré.', 'Installer un fichier de licence.'],
        ['La signature de la licence n\'est pas valide.', 'Fichier modifié ou non émis par l\'éditeur.', 'Demander un nouveau fichier ; ne jamais l\'éditer.'],
        ['Licence expirée.', 'La date d\'expiration est passée : lecture seule.', 'Installer la licence renouvelée.'],
        ['Active users (N) exceed licensed seats (M).', 'Plus de comptes actifs que de postes : lecture seule.', 'Désactiver des comptes ou demander plus de postes.'],
        ['Not included in the licence: …', 'Élément demandé pour une organisation hors licence.', 'Choisir un élément sous licence ou demander une extension.'],
    ], [5.0, 6.0, 6.0], size=9)
    para(doc, 'Les éléments sous licence s\'activent pour chaque organisation dans Administration › Configuration ; seuls ceux de la licence sont acceptés.', size=10)
    shot(doc, 'c-config', 'Figure — Configuration d\'une organisation : packs de fonctionnalités, dans la limite de la licence.', 15)

    h1(doc, 'Données initialisées et comptes')
    seeded_accounts(doc)
    shot(doc, 'c-home', 'Figure — Accueil d\'un projet initialisé, après la première connexion d\'un responsable qualité.', 15)

    h1(doc, 'Exploitation')
    operations(doc, 'dynamicms', 'dynamicms-data')
    h2(doc, 'Sécurité')
    bullets(doc, [
        'Conserver .env avec les droits 600 et ne jamais le diffuser ; effacer ADMIN_PASSWORD et SEED_PASSWORD une fois les mots de passe changés.',
        'Publier l\'application uniquement via le proxy HTTPS ; laisser HTTP_BIND à 127.0.0.1.',
        'Appliquer les mises à jour du système et de Docker ; copier les sauvegardes hors du serveur.',
        'Désactiver les comptes inutilisés (Organisation › Utilisateurs) : ils libèrent des postes de la licence.',
    ], size=10)

    h1(doc, 'Dépannage')
    troubleshooting(doc)
    h2(doc, 'Informations à transmettre au support')
    code(doc, ['sudo docker compose ps', 'sudo docker logs --tail 200 dynamicms', 'curl -s http://127.0.0.1:4000/api/health'])

    path = os.path.join(OUT, f'{G_CUST}.docx')
    doc.save(path)
    print('saved', path)


# ======================================================================= editor guide
def editor():
    brand.set_lang('fr')
    doc = new_document('Guide d\'installation — Runtime éditeur', 'DynamicMS — Guide d\'installation du runtime éditeur (Linux et Windows)')
    cover(doc, 'Guide d\'installation', 'Runtime éditeur', 'DynamicMS 2026.10 — cloud privé ou serveur sur site (Linux) et poste local (Windows), avec la gestion des licences',
          'AI Value · DynamicMS 2026.10 · octobre 2026')
    toc(doc)

    h1(doc, 'À propos de ce guide')
    para(doc, 'Ce guide s\'adresse à l\'équipe de l\'éditeur qui installe sa propre instance de DynamicMS et émet les licences des clients. Deux scénarios sont couverts : un serveur Linux en cloud privé ou sur site (scénario A) et un poste Windows en local (scénario B).')
    bullets(doc, [('Runtime : ', 'l\'application est livrée exécutable, sans code source : serveur minifié et compilé en bytecode V8, interface web minifiée, chargée à la demande et précompressée.'),
                  ('Données initialisées : ', 'tous les secteurs et les PME, chargés au premier démarrage.'),
                  ('Mode production : ', 'aucun compte affiché sur la page de connexion ; mot de passe initial à remplacer à la première connexion.'),
                  ('Gestion des licences : ', 'création, consultation, modification (réémission), révocation, suppression et téléchargement des licences des clients (chapitre 8).')], size=10)

    h1(doc, 'Contenu de la livraison')
    table(doc, ['Fichier', 'Contenu'], [
        [ED_IMG, 'Image du runtime éditeur (conteneur Linux x86_64), utilisée sous Linux et sous Windows (Docker Desktop exécute des conteneurs Linux).'],
        [ED_IMG + '.sha256', 'Empreinte SHA-256 de l\'archive.'],
        [ED_LIN_ZIP, 'Dossier dynamicms-editor (Linux) : docker-compose.yml, .env.example, install.sh, backup.sh, restore.sh, nginx/dynamicms.conf, secrets/LISEZMOI.txt, LISEZMOI.txt.'],
        [ED_WIN_ZIP, 'Dossier dynamicms-editor (Windows) : docker-compose.yml, .env.example, Install-DynamicMS.ps1, Start-DynamicMS.ps1, Stop-DynamicMS.ps1, Backup-DynamicMS.ps1, Restore-DynamicMS.ps1, secrets\\LISEZMOI.txt, LISEZMOI.txt.'],
        [KEY, 'Clé privée de signature des licences (secrète, chapitre 4).'],
        [CUST_IMG + ' (+ .sha256), ' + CUST_ZIP, 'Runtime client à remettre aux clients avec leur licence (voir le guide d\'installation client).'],
        [G_ED + '.docx / .pdf', 'Le présent guide.'],
    ], [7.5, 9.5], size=9)

    h1(doc, 'Le runtime éditeur')
    table(doc, ['', 'Runtime éditeur', 'Runtime client'], [
        ['Usage', 'Instance de l\'éditeur ; émission des licences', 'Installation chez le client'],
        ['Licence requise', 'Non (mode SaaS de l\'éditeur)', 'Oui : lecture seule sans licence valide'],
        ['Gestion des licences', 'Oui (Administration › Licences)', 'Non'],
        ['Éléments commerciaux', 'Visibles (packs, tarifs, devis)', 'Masqués'],
        ['Données initialisées', 'Tous les secteurs et PME', 'Tous les secteurs et PME'],
        ['Comptes de démonstration', 'Aucun', 'Aucun'],
        ['Code source', 'Non livré (bytecode, interface minifiée)', 'Non livré (bytecode, interface minifiée)'],
    ], [3.6, 6.7, 6.7], size=9)
    callout(doc, 'Laisser DEPLOYMENT_MODE à sa valeur par défaut : l\'instance de l\'éditeur fonctionne en mode SaaS, sans licence. Le mode sur site (licence obligatoire) est celui du runtime client.', 'Important —')

    h1(doc, 'La clé de signature des licences')
    bullets(doc, [
        ('Rôle : ', f'{KEY} signe chaque licence. Les runtimes clients contiennent la clé publique correspondante (empreinte {FP}) et n\'acceptent que les licences signées par cette clé.'),
        ('Confidentialité : ', 'quiconque détient la clé privée peut émettre des licences valides. Ne jamais l\'envoyer à un client, la mettre dans un dépôt de code, un e-mail ou une image Docker.'),
        ('Sauvegarde : ', 'en conserver deux copies hors ligne (coffre-fort numérique, support chiffré). Sa perte empêche d\'émettre de nouvelles licences pour les runtimes clients déjà livrés.'),
        ('Contrôle : ', f'Administration › Licences affiche « Clé de signature prête » et l\'empreinte {FP}. Un message rouge signale une clé illisible ou différente.'),
        ('Installation : ', 'la clé est montée en lecture seule dans le conteneur (secrets/licence_private.pem → /run/secrets/licence_private.pem). Sous Linux, elle appartient à l\'utilisateur 1000 avec les droits 400 (install.sh s\'en charge).'),
    ], size=10)

    h1(doc, 'Scénario A — Linux (cloud privé ou serveur sur site)')
    h2(doc, 'Prérequis')
    table(doc, ['Élément', 'Minimum', 'Recommandé'], [
        ['Système', 'Linux x86_64 : Ubuntu 22.04/24.04 LTS, Debian 12, RHEL/Rocky/AlmaLinux 9', 'Ubuntu 24.04 LTS'],
        ['Processeur / mémoire', '2 vCPU / 4 Go', '4 vCPU / 8 Go'],
        ['Disque', '20 Go libres', '50 Go'],
        ['Logiciels', 'Docker Engine 24, Docker Compose v2, openssl, curl, unzip', 'Versions à jour'],
        ['Réseau', 'Port 443 vers le proxy HTTPS ; port 4000 local', 'Nom DNS et certificat TLS'],
    ], [3.4, 8.0, 5.6], size=9)
    h2(doc, 'Installer Docker')
    docker_install(doc)
    h2(doc, 'Installer le runtime éditeur')
    code(doc, ['sudo mkdir -p /opt/dynamicms && cd /opt/dynamicms', f'# copier ici : {ED_IMG}, son .sha256, {ED_LIN_ZIP} et {KEY}', f'sha256sum -c {ED_IMG}.sha256', f'unzip {ED_LIN_ZIP}', 'cd dynamicms-editor', 'sudo ./install.sh'])
    para(doc, 'install.sh contrôle Docker, vérifie l\'empreinte, charge l\'image dynamicms-editor:2026.10, place la clé privée dans secrets/ avec les bons droits, crée .env avec un secret aléatoire, demande l\'administrateur et le mot de passe initial des comptes initialisés, démarre le service et attend l\'application. Il est ensuite possible de supprimer la copie de la clé dans /opt/dynamicms (celle de secrets/ suffit) après l\'avoir sauvegardée hors ligne.')
    h2(doc, 'Paramètres du fichier .env')
    env_table(doc, editor=True)
    h2(doc, 'Publier en HTTPS')
    code(doc, ['sudo cp nginx/dynamicms.conf /etc/nginx/conf.d/dynamicms.conf', '# adapter server_name et les certificats, puis :', 'sudo nginx -t && sudo systemctl reload nginx'])
    para(doc, 'Renseigner CORS_ORIGIN avec l\'adresse publique puis : sudo docker compose up -d.', size=10)

    h1(doc, 'Scénario B — Windows en local (Docker Desktop)')
    h2(doc, 'Prérequis')
    table(doc, ['Élément', 'Exigence'], [
        ['Système', 'Windows 10 22H2 ou Windows 11 (64 bits), virtualisation activée dans le BIOS'],
        ['Docker', 'Docker Desktop 4.x avec le moteur WSL 2 (conteneurs Linux)'],
        ['Mémoire / disque', '8 Go de RAM (4 Go alloués à Docker) ; 20 Go libres'],
        ['Accès', 'Compte administrateur du poste pour installer Docker Desktop'],
        ['Licence Docker Desktop', 'Gratuite pour les petites entreprises, l\'usage personnel et l\'enseignement ; abonnement Docker requis au-delà (conditions de Docker Inc.).'],
    ], [4.0, 13.0], size=9)
    h2(doc, 'Installer Docker Desktop')
    numbered(doc, [
        'Télécharger Docker Desktop pour Windows depuis le site de Docker et lancer l\'installation en cochant « Use WSL 2 instead of Hyper-V ».',
        'Redémarrer le poste si demandé, puis lancer Docker Desktop et attendre « Engine running ».',
        'Vérifier dans PowerShell : docker version et docker compose version.',
    ])
    h2(doc, 'Installer le runtime éditeur')
    numbered(doc, [
        f'Créer un dossier, par exemple C:\\DynamicMS, et y copier {ED_IMG}, son .sha256, {ED_WIN_ZIP} et {KEY}.',
        f'Extraire {ED_WIN_ZIP} dans C:\\DynamicMS (clic droit › Extraire tout) : le dossier dynamicms-editor est créé.',
        'Ouvrir PowerShell dans C:\\DynamicMS\\dynamicms-editor et débloquer les scripts téléchargés : Get-ChildItem *.ps1 | Unblock-File',
        'Lancer l\'installation (voir ci-dessous) et répondre aux questions : e-mail et mot de passe provisoire de l\'administrateur, mot de passe initial des comptes initialisés.',
        'Le navigateur s\'ouvre sur http://localhost:4000 quand l\'application est prête (environ une minute au premier démarrage).',
    ])
    code(doc, ['cd C:\\DynamicMS\\dynamicms-editor', 'powershell -ExecutionPolicy Bypass -File .\\Install-DynamicMS.ps1'])
    h2(doc, 'Utilisation quotidienne')
    table(doc, ['Script', 'Effet'], [
        ['Start-DynamicMS.ps1', 'Démarre l\'application (Docker Desktop doit être lancé).'],
        ['Stop-DynamicMS.ps1', 'Arrête l\'application ; les données sont conservées.'],
        ['Backup-DynamicMS.ps1', 'Sauvegarde complète dans .\\sauvegardes (arrêt de quelques secondes).'],
        ['Restore-DynamicMS.ps1 -Archive <fichier>', 'Restaure une sauvegarde (confirmation RESTAURER).'],
    ], [6.0, 11.0], size=9)
    para(doc, 'Sous Windows, l\'application n\'est accessible que depuis le poste (127.0.0.1). Limiter l\'accès au dossier secrets à votre compte (Propriétés › Sécurité).', size=10)

    h1(doc, 'Premier accès')
    para(doc, 'Ouvrir l\'adresse de l\'application : la page de connexion n\'affiche aucun compte. Se connecter avec ADMIN_EMAIL et le mot de passe provisoire, puis choisir un nouveau mot de passe.')
    shot(doc, 'c-login', 'Figure — Page de connexion : aucun compte de démonstration n\'est affiché.', 14)

    h1(doc, 'Gestion des licences')
    para(doc, 'Administration › onglet Licences (administrateur de la plateforme uniquement). L\'écran affiche l\'état de la clé de signature et le registre des licences des clients.')
    shot(doc, 'e-licences-empty', f'Figure — Clé de signature prête (empreinte {FP}) et registre encore vide.', 15)
    h2(doc, 'Créer une licence')
    numbered(doc, [
        'Sélectionner Nouvelle licence.',
        'Renseigner le client, le contact chez le client, les postes (comptes actifs autorisés), la date d\'expiration et le mode de déploiement (DEP-4 cloud privé du client, DEP-5 centre de données du client, DEP-6 hybride, DEP-7 isolé).',
        'Cocher les éléments sous licence : au moins un pack de base, puis les packs sectoriels et de capacités, les options et les normes de conformité et de sécurité utiles.',
        'Sélectionner Émettre la licence : la licence reçoit un code LIC-AAAA-NNN, la version 1 et le statut Actif.',
    ])
    shot(doc, 'e-licence-form', 'Figure — Formulaire de création d\'une licence.', 15)
    table(doc, ['Champ', 'Règle'], [
        ['Client', 'Obligatoire.'],
        ['Postes', 'Nombre entier ≥ 1 : comptes actifs autorisés, administrateur de la plateforme exclu. Une installation client avec données initialisées en compte 1 140.'],
        ['Expire le', 'Date postérieure à aujourd\'hui. Après cette date, l\'installation client passe en lecture seule.'],
        ['Mode de déploiement', 'DEP-4 à DEP-7 (modes sur site).'],
        ['Éléments sous licence', 'Au moins un pack de base (DMS-SME, DMS-PRO ou DMS-ENT).'],
    ], [4.0, 13.0], size=9)
    h2(doc, 'Consulter et télécharger')
    para(doc, 'Le registre liste chaque licence avec son code, son client, ses postes, son expiration, ses éléments, son mode, sa version et son statut (Actif, Expirée, Révoquée). L\'icône de téléchargement produit le fichier Licence_<code>_v<version>_<client>.lic à remettre au client ; Télécharger la clé publique produit la clé de vérification (déjà intégrée aux runtimes clients).')
    shot(doc, 'e-licences-list', 'Figure — Registre des licences après l\'émission de LIC-2026-001.', 15)
    h2(doc, 'Modifier (réémettre)')
    para(doc, 'L\'icône de modification ouvre la licence ; une raison est obligatoire. Chaque modification émet une nouvelle version signée (v2, v3…) ; l\'historique des versions est conservé. Le client installe le nouveau fichier, qui remplace l\'ancien.')
    shot(doc, 'e-licence-edit', 'Figure — Modification d\'une licence : passage à 2 000 postes, avec sa raison.', 15)
    shot(doc, 'e-licences-v2', 'Figure — La licence passe en version 2.', 15)
    h2(doc, 'Révoquer')
    para(doc, 'L\'icône de révocation demande une raison. Une licence révoquée ne peut plus être téléchargée ni modifiée. Les installations clientes fonctionnant sans connexion à l\'éditeur, une licence déjà installée reste valable jusqu\'à sa date d\'expiration ou jusqu\'à l\'installation d\'une autre licence : prévoir des durées de licence courtes (un an) et les renouveler.')
    shot(doc, 'e-licence-revoke', 'Figure — Révocation d\'une licence.', 15)
    h2(doc, 'Supprimer')
    para(doc, 'La suppression retire la licence et son historique du registre ; la piste d\'audit garde la trace. Une licence déjà remise et encore valide doit d\'abord être révoquée.')
    h2(doc, 'Traçabilité')
    para(doc, 'Chaque émission, modification, téléchargement, révocation et suppression est enregistré dans Administration › Piste d\'audit, avec l\'auteur, la date et les valeurs avant et après.')
    h2(doc, 'Livrer un client')
    numbered(doc, [
        f'Remettre au client : {CUST_IMG} et son .sha256, {CUST_ZIP}, le guide d\'installation client et le guide utilisateur.',
        'Émettre sa licence (postes, date, éléments) et lui envoyer le fichier .lic par un canal distinct.',
        'Le client installe le runtime puis la licence (Administration › Exploitation › Licence).',
        'Au renouvellement ou à l\'extension, modifier la licence et envoyer la nouvelle version.',
    ])

    h1(doc, 'Données initialisées et comptes')
    seeded_accounts(doc)

    h1(doc, 'Exploitation')
    operations(doc, 'dynamicms-editor', 'dynamicms-editor-data', editor=True)
    para(doc, 'Sous Windows, utiliser les scripts Start, Stop, Backup et Restore du dossier (scénario B). Inclure la clé privée dans la sauvegarde hors ligne, séparément des données.', size=10)

    h1(doc, 'Dépannage')
    troubleshooting(doc, editor=True)
    table(doc, ['Symptôme (Windows)', 'Action'], [
        ['« Docker Desktop n\'est pas démarré »', 'Lancer Docker Desktop et attendre « Engine running ».'],
        ['Exécution de scripts désactivée', 'Utiliser powershell -ExecutionPolicy Bypass -File … et Unblock-File.'],
        ['localhost:4000 ne répond pas', 'docker logs dynamicms-editor ; vérifier qu\'aucune autre application n\'utilise le port.'],
    ], [6.0, 11.0], size=9)

    path = os.path.join(OUT, f'{G_ED}.docx')
    doc.save(path)
    print('saved', path)


if __name__ == '__main__':
    which = sys.argv[1] if len(sys.argv) > 1 else 'all'
    if which in ('customer', 'all'):
        customer()
    if which in ('editor', 'all'):
        editor()
