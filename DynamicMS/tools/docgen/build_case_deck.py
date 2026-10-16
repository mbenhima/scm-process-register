"""Builds the universal case-study deck (PowerPoint 16:9, POWERACT chart) for a healthcare
audience, in English or French: one full real run of the 12 end-to-end processes of project
HZ-UNI-QMS, with for each phase its macro processes, one macro process drilled down to its tasks
and steps with the values recorded, and screenshots of the application.
Data: deliverables/build/HZ-UNI-QMS(.fr).json; screenshots: deliverables/build/case/<lang>/.
Usage: python3 build_case_deck.py en|fr"""
import json
import os
import sys

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

import build_deck
from build_deck import Deck, C, rgb, BUILD, ROOT
import build_flow_guide as G

LANG = sys.argv[1] if len(sys.argv) > 1 else 'en'
SHOTS = os.path.join(BUILD, 'case', LANG)
FLOW = os.path.join(BUILD, f'flow-{LANG}')


def X(en, fr):
    return fr if LANG == 'fr' else en


PDCA_NAME = {'P': X('Plan', 'Planifier'), 'D': X('Do', 'Réaliser'), 'C': X('Check', 'Vérifier'), 'A': X('Act', 'Agir')}


def cut(v, n):
    s = '' if v is None else str(v).replace('\n', ' ')
    return s if len(s) <= n else s[:n - 1].rstrip() + '…'


def recorded(step):
    """A short summary of what was recorded in a step."""
    tabs = [f for f in step['type_'] if f.get('columns')]
    simple = [f for f in step['type_'] if not f.get('columns') and f.get('value') not in (None, '')]
    if tabs:
        f = tabs[0]; n = len(f.get('rows') or [])
        first = (f.get('rows') or [[]])[0]
        lead = cut(first[0] if first else '', 38)
        return X(f"{n} rows in “{f['field']}” — {lead}", f"{n} lignes dans « {f['field']} » — {lead}")
    if simple:
        f = simple[0]
        return f"{f['field']}: {cut(f['value'], 60)}" if LANG == 'en' else f"{f['field']} : {cut(f['value'], 60)}"
    return X('Completed (no input)', 'Terminée (sans saisie)')


class CaseDeck(Deck):
    def table(self, s, x, y, w, headers, rows, widths, size=10, row_h=0.3):
        shp = s.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(row_h * (len(rows) + 1)))
        t = shp.table
        tot = sum(widths)
        for j, wd in enumerate(widths):
            t.columns[j].width = Inches(w * wd / tot)
        for i in range(len(rows) + 1):
            t.rows[i].height = Inches(row_h)
        for i, row in enumerate([headers] + rows):
            for j, v in enumerate(row):
                c = t.cell(i, j)
                c.fill.solid(); c.fill.fore_color.rgb = rgb(C['orange'] if i == 0 else (C['white'] if i % 2 else C['light']))
                c.margin_left = c.margin_right = Inches(0.06); c.margin_top = c.margin_bottom = Inches(0.03)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = c.text_frame; tf.word_wrap = True
                p = tf.paragraphs[0]; p.text = ''
                r = p.add_run(); r.text = str(v)
                r.font.size = Pt(size); r.font.name = 'Calibri'
                r.font.bold = i == 0 or (j == 0)
                r.font.color.rgb = rgb(C['white'] if i == 0 else (C['dark'] if j == 0 else C['ink']))
        return shp

    def divider(self, title, sub):
        from pptx.enum.shapes import MSO_SHAPE
        s = self.slide(dark=True); self.section = title
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.8), Inches(1.4), Inches(4.8), Inches(4.8))
        c.fill.background(); c.line.color.rgb = rgb(C['orange']); c.line.width = Pt(6)
        self.text(s, 0.9, 1.9, 8.6, 1.5, title, 36, True, 'white', font='Cambria', anchor=MSO_ANCHOR.BOTTOM)
        self.text(s, 0.9, 3.6, 8.2, 1.0, sub, 18, False, 'line')
        self.footer(s, dark=True)

    def kpi(self, s, x, y, w, value, label, fill='white'):
        self.card(s, x, y, w, 1.15, fill=fill)
        self.text(s, x + 0.2, y + 0.1, w - 0.3, 0.55, value, 26, True, 'deep', font='Cambria')
        self.text(s, x + 0.2, y + 0.68, w - 0.3, 0.42, label, 11, False, 'ink')

    def bullets_card(self, s, x, y, w, h, head, lines, fill='white', size=12):
        self.card(s, x, y, w, h, fill=fill)
        self.text(s, x + 0.22, y + 0.15, w - 0.4, 0.35, head.upper(), 10, True, 'deep', spacing=120)
        self.text(s, x + 0.22, y + 0.5, w - 0.4, h - 0.6, lines, size, False, 'dark')


def main():
    G.LANG = LANG
    data = json.load(open(os.path.join(BUILD, 'HZ-UNI-QMS.json' if LANG == 'en' else 'HZ-UNI-QMS.fr.json'), encoding='utf8'))
    picks = {p['e2e']: p for p in json.load(open(os.path.join(BUILD, 'case', 'picks.json')))['picks']}
    build_deck.TXT[LANG]['date'] = X('POWERACT Consulting · October 2026', 'POWERACT Consulting · octobre 2026')
    build_deck.TXT[LANG]['deck'] = X('Universal Case Study — One Full Run of DynamicMS', 'Étude de cas universelle — Un déroulement complet de DynamicMS')
    build_deck.TXT[LANG]['sub'] = X('Twelve end-to-end processes, their macro processes, tasks and steps, as recorded in the application — prepared for healthcare organizations',
                                    'Douze processus de bout en bout, leurs macro-processus, tâches et étapes, tels qu’enregistrés dans l’application — préparé pour les établissements de santé')
    d = CaseDeck(LANG)
    phases = data['phases']; org = data['org']; prj = data['project']
    n_mps = sum(len(ph['mps']) for ph in phases)
    n_tasks = sum(len(m['tasks']) for ph in phases for m in ph['mps'])
    n_steps = sum(len(t['steps']) for ph in phases for m in ph['mps'] for t in m['tasks'])
    feeds = G.feeds()

    # ---------------------------------------------------------------- opening
    d.cover()
    d.agenda_items = None
    s = d.slide(); d.section = X('Agenda', 'Sommaire')
    d.eyebrow(s, 'DynamicMS'); d.title(s, X('Agenda', 'Sommaire'))
    items = [X('The case: one organization, one full run', 'Le cas : une organisation, un déroulement complet'),
             X('What it means for a healthcare organization', 'Ce que cela signifie pour un établissement de santé'),
             X('The twelve end-to-end processes, phase by phase', 'Les douze processus de bout en bout, phase par phase'),
             X('Evidence, measurement, documents and AI', 'Preuves, mesure, documents et IA'),
             X('Next steps for your organization', 'Prochaines étapes pour votre établissement')]
    for i, it in enumerate(items):
        y = 2.1 + i * 0.85
        d.text(s, 0.9, y, 0.9, 0.7, f'{i + 1:02d}', 30, True, 'deep', font='Cambria')
        d.text(s, 2.0, y + 0.12, 9.5, 0.6, it, 20, False, 'dark')
    d.footer(s)

    # The case
    s = d.slide(); d.section = X('The case', 'Le cas')
    d.eyebrow(s, X('The case', 'Le cas'))
    d.title(s, X(f"{org['name']}: a full QMS run from context to innovation", f"{org['name']} : un déroulement SMQ complet, du contexte à l’innovation"), size=21)
    d.subtitle(s, X(f"Project {prj['code']} · ISO 9001 · {prj['start']} to {prj['end']} · {org['employees']:,} employees · universal profile",
                    f"Projet {prj['code']} · ISO 9001 · du {prj['start']} au {prj['end']} · {org['employees']:,} salariés · profil universel".replace(',', ' ')))
    k = [('12', X('end-to-end processes, run as phases', 'processus de bout en bout, déroulés en phases')),
         (str(n_mps), X('macro processes activated', 'macro-processus activés')),
         (str(n_tasks), X('tasks', 'tâches')),
         (f"{n_steps:,}".replace(',', ' ' if LANG == 'fr' else ','), X('workflow steps with their recorded values', 'étapes de workflow avec leurs valeurs')),
         (f"{prj['progress']}%", X('of the run completed; 11 gates passed (Go)', 'du déroulement terminé ; 11 jalons franchis (Go)')),
         (str(len(data['users'])), X('people, each in a named role', 'personnes, chacune dans un rôle nommé'))]
    for i, (v, lab) in enumerate(k):
        col, row = i % 3, i // 3
        d.kpi(s, 0.7 + col * 4.05, 2.45 + row * 1.45, 3.8, v, lab, fill='tint' if i == 0 else 'white')
    d.caption(s, X('All values come from the demonstration database of DynamicMS; names and figures are fictional.',
                   'Toutes les valeurs proviennent de la base de démonstration de DynamicMS ; noms et chiffres sont fictifs.'), y=5.5)
    d.text(s, 0.7, 5.95, 11.9, 0.6, X(f"Licence of the case: {org['pack']} with {', '.join(org['caps'])} and add-ons {', '.join(org['addons'])}.",
                                      f"Licence du cas : {org['pack']} avec {', '.join(org['caps'])} et options {', '.join(org['addons'])}."), 12, False, 'ink')
    d.footer(s)

    # Healthcare relevance
    s = d.slide(); d.section = X('Healthcare', 'Santé')
    d.eyebrow(s, X('For healthcare', 'Pour la santé'))
    d.title(s, X('The same twelve processes run a hospital’s management system', 'Les mêmes douze processus pilotent le système de management d’un établissement de santé'), size=21)
    d.subtitle(s, X('The universal run is the common core. For a healthcare provider (segment V-29), DynamicMS activates the following on top of it.',
                    'Le déroulement universel est le socle commun. Pour un établissement de soins (segment V-29), DynamicMS active en plus les éléments suivants.'))
    d.bullets_card(s, 0.7, 2.4, 3.8, 3.9, X('Standards', 'Normes'),
                   ['JCI', 'ISO 9001', 'ISO/IEC 27001', 'HIPAA · GDPR', X('National Health Authority', 'Autorité nationale de santé'), 'HL7 FHIR'], fill='tint')
    d.bullets_card(s, 4.75, 2.4, 3.8, 3.9, X('79 macro processes, including', '79 macro-processus, dont'),
                   [X('Medical Laboratory Management', 'Gestion du laboratoire médical'), X('Facility Management', 'Gestion des installations'),
                    X('IT Service Management', 'Gestion des services informatiques'), X('Cloud Security and Privacy', 'Sécurité du cloud et vie privée'),
                    X('Anti-Bribery and Corruption', 'Lutte contre la corruption')])
    d.bullets_card(s, 8.8, 2.4, 3.85, 3.9, X('Recommended licence', 'Licence recommandée'),
                   ['DMS-PRO', X('+ DMS-LSCI Life Sciences & Healthcare', '+ DMS-LSCI Sciences de la vie et santé'), X('+ DMS-CERT Certification', '+ DMS-CERT Certification'), X('+ DMS-AI Intelligence & Assurance', '+ DMS-AI Intelligence et assurance')])
    d.caption(s, X('Source: DynamicMS catalog (segment V-29 Healthcare Providers, Packs catalog 2.5).', 'Source : catalogue DynamicMS (segment V-29 Santé — établissements de soins, catalogue des packs 2.5).'), y=6.5)
    d.footer(s)

    # Reading the run in a hospital
    s = d.slide(); d.section = X('Healthcare', 'Santé')
    d.eyebrow(s, X('For healthcare', 'Pour la santé'))
    d.title(s, X('How the run reads in a hospital', 'Comment le déroulement se lit dans un hôpital'), size=26)
    rows = [
        ['E2E-01', X('Interested parties, context issues', 'Parties intéressées, enjeux du contexte'), X('Patients and families, regulator, medical staff, insurers', 'Patients et familles, autorité de tutelle, personnel médical, assureurs')],
        ['E2E-03', X('Risk and opportunity register', 'Registre des risques et opportunités'), X('Patient safety risks, infection control, continuity of care', 'Risques pour la sécurité des patients, hygiène, continuité des soins')],
        ['E2E-05', X('Competence, documents, measuring equipment', 'Compétences, documents, équipements de mesure'), X('Clinical competences, care protocols, calibrated medical devices', 'Compétences cliniques, protocoles de soins, dispositifs médicaux étalonnés')],
        ['E2E-08', X('Requirements to delivery', 'Des exigences à la livraison'), X('Care pathways from admission to discharge', 'Parcours de soins, de l’admission à la sortie')],
        ['E2E-09', X('Nonconformity, root cause, corrective action', 'Non-conformité, cause racine, action corrective'), X('Adverse events and near misses, analysis, lessons learned', 'Événements indésirables et presque-accidents, analyse, retours d’expérience')],
        ['E2E-10', X('KPIs, process and management reviews', 'KPI, revues de processus et de direction'), X('Quality and safety indicators, accreditation readiness', 'Indicateurs qualité et sécurité, préparation à l’accréditation')],
        ['E2E-11', X('Audits and assurance', 'Audits et assurance'), X('Internal audits against JCI and ISO 9001, mock surveys', 'Audits internes JCI et ISO 9001, visites blanches')],
    ]
    d.table(s, 0.7, 1.95, 11.9, ['E2E', X('In the universal run', 'Dans le déroulement universel'), X('In a hospital (illustration)', 'Dans un hôpital (illustration)')], rows, [1.1, 4.6, 6.2], size=12, row_h=0.52)
    d.caption(s, X('Illustration of how the same steps are used; the healthcare macro processes add their own steps.', 'Illustration de l’usage des mêmes étapes ; les macro-processus propres à la santé ajoutent leurs étapes.'), y=6.4)
    d.footer(s)

    # PDCA + lifecycle
    s = d.slide(); d.section = X('The run', 'Le déroulement')
    d.eyebrow(s, 'PDCA'); d.title(s, X('Twelve phases on one Plan-Do-Check-Act cycle', 'Douze phases sur un cycle Planifier-Réaliser-Vérifier-Agir'), size=26)
    d.picture(s, os.path.join(BUILD, f'pdca-{LANG}.png'), 0.9, 1.9, 4.6, 4.75)
    rows = [[ph['id'], ph['name'], PDCA_NAME[G.PDCA[ph['id']]], str(len(ph['mps'])), str(sum(len(t['steps']) for m in ph['mps'] for t in m['tasks'])),
             X('Go', 'Go') if ph.get('decision') == 'Go' else X('Active', 'En cours')] for ph in phases]
    d.table(s, 6.0, 1.85, 6.6, ['E2E', X('Name', 'Nom'), 'PDCA', 'MP', X('Steps', 'Étapes'), X('Gate', 'Jalon')], rows, [0.9, 3.2, 1.1, 0.6, 0.8, 0.8], size=10, row_h=0.36)
    d.footer(s)

    for name, title, cap in (('home', X('The project home: progress, next steps, KPIs', 'L’accueil du projet : avancement, prochaines étapes, KPI'), X('Home of project HZ-UNI-QMS.', 'Accueil du projet HZ-UNI-QMS.')),
                             ('lifecycle', X('The lifecycle: twelve phases and their gates', 'Le cycle de vie : douze phases et leurs jalons'), X('Lifecycle screen: each card is an end-to-end process with its progress and gate.', 'Écran Cycle de vie : chaque carte est un processus de bout en bout avec son avancement et son jalon.'))):
        s = d.slide(); d.section = X('The run', 'Le déroulement')
        d.eyebrow(s, X('Screens', 'Écrans')); d.title(s, title, size=26)
        d.picture(s, os.path.join(SHOTS, f'{name}.png'), 0.9, 1.75, 11.5, 4.75)
        d.caption(s, cap, y=6.62)
        d.footer(s)

    d.divider(X('The full run, phase by phase', 'Le déroulement complet, phase par phase'),
              X('For each end-to-end process: its macro processes, one macro process down to its steps, and the screens.', 'Pour chaque processus de bout en bout : ses macro-processus, un macro-processus jusqu’à ses étapes, et les écrans.'))

    # ---------------------------------------------------------------- 12 phases
    for i, ph in enumerate(phases):
        e = ph['id']; k = G.PDCA[e]; pk = picks[e]
        mps = ph['mps']
        tasks = sum(len(m['tasks']) for m in mps); steps = sum(len(t['steps']) for m in mps for t in m['tasks'])
        sec = f"{e} — {ph['name']}"
        # A. overview
        s = d.slide(); d.section = sec
        d.eyebrow(s, f"{X('Phase', 'Phase')} {i + 1:02d} · {PDCA_NAME[k]}")
        d.title(s, sec, size=26)
        d.text(s, 0.7, 1.7, 5.4, 0.9, ph['goals'] + '.', 14, False, 'ink')
        for j, (v, lab) in enumerate(((str(len(mps)), X('macro processes', 'macro-processus')), (str(tasks), X('tasks', 'tâches')), (str(steps), X('steps', 'étapes')))):
            d.kpi(s, 0.7 + j * 1.85, 2.6, 1.7, v, lab)
        d.bullets_card(s, 0.7, 3.95, 5.4, 2.6, X('From trigger to result', 'Du déclencheur au résultat'),
                       [X('Trigger: ', 'Déclencheur : ') + cut(ph['trigger'], 110), X('Result: ', 'Résultat : ') + cut(ph['terminal'], 110),
                        X('Gate: ', 'Jalon : ') + (ph['gate']['name'] if ph.get('gate') else '—') + (f" — {ph['decision']}" if ph.get('decision') else X(' — active', ' — en cours'))], fill='tint', size=11)
        half = (len(mps) + 1) // 2 if len(mps) > 9 else len(mps)
        cols = [mps[:half], mps[half:]] if len(mps) > 9 else [mps]
        cw = 3.15 if len(cols) == 2 else 6.3
        d.text(s, 6.5, 1.72, 6.2, 0.3, X('MACRO PROCESSES OF THIS PHASE', 'MACRO-PROCESSUS DE CETTE PHASE'), 10, True, 'deep', spacing=120)
        for ci, col in enumerate(cols):
            lines = [f"{m['id']}  {cut(m['name'], 44 if len(cols) == 2 else 70)}" for m in col]
            d.text(s, 6.5 + ci * cw, 2.05, cw - 0.1, 4.6, lines, 10 if len(cols) == 2 else 11, False, 'dark')
        d.footer(s)

        # B. phase screen
        s = d.slide(); d.section = sec
        d.eyebrow(s, f"{e} · {X('Screen', 'Écran')}"); d.title(s, X(f'{ph["name"]}: the phase in the application', f'{ph["name"]} : la phase dans l’application'), size=24)
        d.picture(s, os.path.join(SHOTS, f'phase-{e}.png'), 0.78, 1.85, 8.35, 4.7)
        d.bullets_card(s, 9.5, 1.85, 3.15, 2.2, X('Delivers', 'Produit'), [cut(feeds[e][0], 150)], size=11)
        d.bullets_card(s, 9.5, 4.2, 3.15, 2.35, X('Feeds', 'Alimente'), [cut(feeds[e][1], 150)], fill='tint', size=11)
        d.caption(s, X(f'Lifecycle › {e}: macro processes of the phase with their progress and gate.', f'Cycle de vie › {e} : macro-processus de la phase avec leur avancement et leur jalon.'), y=6.62, w=8.4)
        d.footer(s)

        # C. macro process drill-down
        mp = next(m for m in mps if m['id'] == pk['mp'])
        s = d.slide(); d.section = sec
        d.eyebrow(s, f"{X('Macro process', 'Macro-processus')} · {mp['id']} ({mp['code']})")
        d.title(s, mp['name'], size=24)
        d.text(s, 0.7, 1.62, 12, 0.4, f"{mp['goal']}  ·  {X('Owner', 'Pilote')}: {mp['owner']}" + (f"  ·  {mp['clauses']}" if mp.get('clauses') else ''), 12, False, 'ink', italic=True)
        all_steps = [(t, st) for t in mp['tasks'] for st in t['steps']]
        rows = [[st['id'], cut(st['name'], 40), cut(st['role'], 22), cut(recorded(st), 70)] for t, st in all_steps]
        if len(rows) > 13:
            rows = rows[:12] + [['…', X(f'{len(all_steps) - 12} more steps', f'{len(all_steps) - 12} autres étapes'), '', '']]
        rh = 0.34 if len(rows) <= 11 else 0.3
        d.table(s, 0.7, 2.1, 7.7, [X('Step', 'Étape'), X('Name', 'Nom'), X('Role', 'Rôle'), X('What was recorded', 'Ce qui a été saisi')], rows, [0.9, 2.6, 1.5, 3.4], size=9, row_h=rh)
        d.picture(s, os.path.join(SHOTS, f'mp-{e}.png'), 8.75, 2.15, 3.9, 3.0)
        d.text(s, 8.7, 5.35, 4.0, 1.2, [X(f"{len(mp['tasks'])} tasks · {len(all_steps)} steps", f"{len(mp['tasks'])} tâches · {len(all_steps)} étapes")] +
               [f"{t['id'].split('.')[-1]}  {cut(t['name'], 42)}" for t in mp['tasks'][:4]], 10, False, 'dark')
        d.footer(s)

        # D. step screen
        st = next(x for t in mp['tasks'] for x in t['steps'] if x['id'] == pk['step'])
        tab = next(f for f in st['type_'] if f.get('columns'))
        s = d.slide(); d.section = sec
        d.eyebrow(s, f"{X('Step', 'Étape')} · {st['id']} · {st['role']}")
        d.title(s, st['name'], size=24)
        d.picture(s, os.path.join(SHOTS, f'step-{e}.png'), 0.78, 1.85, 8.35, 4.7)
        lines = [f"• {cut(r[0], 34)}" + (f" — {cut(r[1], 26)}" if len(r) > 1 and r[1] else '') for r in (tab.get('rows') or [])[:6]]
        d.bullets_card(s, 9.5, 1.85, 3.15, 4.7, X('What was recorded', 'Ce qui a été saisi'),
                       [cut(tab['field'], 60), ''] + lines, size=10)
        d.caption(s, X(f"Step {st['id']} ({st['form']}): rows typed, pasted or imported, saved as they are entered.", f"Étape {st['id']} ({st['form']}) : lignes saisies, collées ou importées, enregistrées au fil de la saisie."), y=6.62, w=8.4)
        d.footer(s)

    # ---------------------------------------------------------------- evidence and closing
    d.divider(X('Evidence, measurement and AI', 'Preuves, mesure et IA'), X('What the run produces for audits, reviews and accreditation.', 'Ce que le déroulement produit pour les audits, les revues et l’accréditation.'))
    feats = [
        ('nc', X('Nonconformity', 'Non-conformité'), X('From detection to lessons learned', 'De la détection au retour d’expérience'), [X('Containment, root cause, corrective actions', 'Confinement, cause racine, actions correctives'), X('Effectiveness checked by another person', 'Efficacité vérifiée par une autre personne'), X('Cost recorded with its currency', 'Coût enregistré avec sa devise')]),
        ('audit', X('Audit', 'Audit'), X('Audits with findings and conclusions', 'Audits avec constats et conclusions'), [X('Programme, checklist, findings', 'Programme, liste de contrôle, constats'), X('Findings become nonconformities', 'Les constats deviennent des non-conformités'), X('Ready for JCI and ISO surveys', 'Prêt pour les visites JCI et ISO')]),
        ('kpis', 'KPI', X('Indicators against their targets', 'Indicateurs face à leurs cibles'), [X('Monthly values and trend', 'Valeurs mensuelles et tendance'), X('Data table, CSV and image export', 'Tableau de données, export CSV et image'), X('Feed the management review', 'Alimentent la revue de direction')]),
        ('risks', X('Risks', 'Risques'), X('Risks and opportunities by score', 'Risques et opportunités par cotation'), [X('Likelihood × impact heatmap', 'Carte de chaleur probabilité × impact'), X('Treatment and owner', 'Traitement et responsable'), X('Linked to the steps they govern', 'Liés aux étapes qu’ils encadrent')]),
        ('documents', X('Documents', 'Documents'), X('Controlled documents generated from the data', 'Documents maîtrisés générés à partir des données'), [X('37 IMS templates', '37 modèles SMI'), X('Versions, review and approval', 'Versions, revue et approbation'), X('PDF, Word, Excel, PowerPoint', 'PDF, Word, Excel, PowerPoint')]),
        ('reports', X('Reports', 'Rapports'), X('Reports for every review', 'Des rapports pour chaque revue'), [X('Management review inputs (ISO 9001 §9.3.2)', 'Éléments d’entrée de la revue de direction (ISO 9001 §9.3.2)'), X('English, French, Arabic', 'Anglais, français, arabe'), X('PDF, Excel, Word, PowerPoint, CSV', 'PDF, Excel, Word, PowerPoint, CSV')]),
        ('racsi', 'RACSI', X('Who is responsible for each step', 'Qui est responsable de chaque étape'), [X('Responsible, Accountable, Support, Consulted, Informed', 'Responsable, Autorité, Soutien, Consulté, Informé'), X('Linked to the organization chart', 'Lié à l’organigramme'), X('Exported for audits', 'Exportée pour les audits')]),
    ]
    for shot, eb, title, bl in feats:
        p = os.path.join(SHOTS, f'{shot}.png')
        if not os.path.exists(p):
            continue
        s = d.slide(); d.section = X('Evidence', 'Preuves')
        d.eyebrow(s, eb); d.title(s, title, size=26)
        d.picture(s, p, 0.78, 1.95, 8.35, 4.62)
        for j, b in enumerate(bl):
            y = 2.05 + j * 1.45
            d.card(s, 9.55, y, 3.1, 1.25)
            d.text(s, 9.75, y + 0.18, 2.75, 0.95, b, 14, False, 'dark', anchor=MSO_ANCHOR.MIDDLE)
        d.footer(s)
    # AI
    s = d.slide(); d.section = X('Evidence', 'Preuves')
    d.eyebrow(s, X('Governed AI', 'IA gouvernée')); d.title(s, X('AI proposes, a person decides', 'L’IA propose, une personne décide'), size=26)
    d.picture(s, os.path.join(FLOW, 'ai-proposal.png'), 0.78, 1.95, 8.35, 4.62)
    for j, b in enumerate([X('Each suggestion comes from an approved AI use case', 'Chaque suggestion provient d’un cas d’usage IA approuvé'),
                           X('Each item accepted, edited or rejected', 'Chaque élément accepté, modifié ou rejeté'),
                           X('Every decision kept in an append-only log', 'Chaque décision conservée dans un journal non modifiable')]):
        y = 2.05 + j * 1.45
        d.card(s, 9.55, y, 3.1, 1.25)
        d.text(s, 9.75, y + 0.18, 2.75, 0.95, b, 14, False, 'dark', anchor=MSO_ANCHOR.MIDDLE)
    d.footer(s)

    # Next steps
    s = d.slide(); d.section = X('Next steps', 'Prochaines étapes')
    d.eyebrow(s, X('Next steps', 'Prochaines étapes')); d.title(s, X('From this case to your organization', 'De ce cas à votre établissement'), size=28)
    steps_ = [X('Scoping workshop: your sites, services, standards (JCI, ISO 9001, ISO 15189 for the laboratory) and priorities.', 'Atelier de cadrage : vos sites, services, normes (JCI, ISO 9001, ISO 15189 pour le laboratoire) et priorités.'),
              X('Configuration: healthcare segment V-29, licence, organization chart, roles and users.', 'Configuration : segment santé V-29, licence, organigramme, rôles et utilisateurs.'),
              X('Pilot: run E2E-01 to E2E-04 on one site with your data; first gate decision.', 'Pilote : dérouler E2E-01 à E2E-04 sur un site avec vos données ; première décision de jalon.'),
              X('Deployment: the twelve phases across services, with reviews and reports in French and English.', 'Déploiement : les douze phases dans tous les services, avec revues et rapports en français et en anglais.')]
    for j, line in enumerate(steps_):
        y = 2.0 + j * 1.1
        d.card(s, 0.7, y, 11.9, 0.9)
        d.text(s, 0.95, y + 0.12, 0.8, 0.65, f'{j + 1:02d}', 24, True, 'deep', font='Cambria', anchor=MSO_ANCHOR.MIDDLE)
        d.text(s, 1.85, y + 0.17, 10.5, 0.6, line, 15, False, 'dark', anchor=MSO_ANCHOR.MIDDLE)
    d.text(s, 0.7, 6.45, 11.9, 0.4, X('The full step-by-step values of this run are in the DynamicMS user guide (full real run).', 'Les valeurs détaillées de chaque étape de ce déroulement figurent dans le guide utilisateur DynamicMS (déroulement réel complet).'), 12, False, 'ink', italic=True)
    d.footer(s)

    out = os.path.join(ROOT, 'deliverables', f'DynamicMS_Case_Study_Healthcare_{LANG.upper()}.pptx')
    d.prs.save(out)
    print('saved', out, d.n, 'slides')


if __name__ == '__main__':
    main()
