#!/usr/bin/env python3
"""Extract the DynamicMS source documents (docs/sources) into the English
catalog used by the server seed (server/seed/catalog/catalog.en.json).

Run:  python3 tools/extract_sources.py
Requires: python-docx, openpyxl
"""
import json, os, re, sys
import docx, openpyxl
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'docs', 'sources')
OUT = os.path.join(ROOT, 'server', 'seed', 'catalog')
os.makedirs(OUT, exist_ok=True)


def blocks(path):
    d = docx.Document(path)
    out = []
    for child in d.element.body.iterchildren():
        if child.tag.endswith('}p'):
            p = Paragraph(child, d)
            t = p.text.strip()
            if not t:
                continue
            s = p.style.name if p.style is not None else ''
            lvl = 0
            if s.startswith('Heading') and s.split()[-1].isdigit():
                lvl = int(s.split()[-1])
            out.append(('p', lvl, t))
        elif child.tag.endswith('}tbl'):
            tb = Table(child, d)
            rows = []
            for r in tb.rows:
                cells, prev = [], None
                for c in r.cells:
                    if c._tc is prev:
                        continue
                    prev = c._tc
                    cells.append(c.text.strip())
                rows.append(cells)
            out.append(('t', 0, rows))
    return out


def split_list(s):
    return [x.strip() for x in re.split(r',\s*(?![^()]*\))', s) if x.strip()]


# ---------------------------------------------------------------- Process design
pd = blocks(os.path.join(SRC, '01_DynamicMS_Process_Design_E2E_v7.docx'))
cat = {}

def find_table_after(pred, start=0):
    for i in range(start, len(pd)):
        k, lvl, v = pd[i]
        if k == 'p' and pred(v):
            for j in range(i + 1, len(pd)):
                if pd[j][0] == 't':
                    return pd[j][2], j
    return None, -1

# E2E catalogue
tbl, _ = find_table_after(lambda v: v.startswith('1.3 End-to-End Process Catalogue'))
e2e = []
for r in tbl[1:]:
    e2e.append({'id': r[0], 'type': r[1], 'name': r[2], 'goals': r[3], 'ufCount': r[5]})
cat['e2e'] = e2e

tbl, _ = find_table_after(lambda v: v.startswith('1.4 Root, Partition'))
cat['e2eTypes'] = [{'type': r[0], 'definition': r[1], 'members': r[2]} for r in tbl[1:]]

# Segments overview
tbl, _ = find_table_after(lambda v: v.startswith('2.1 Segments Overview'))
segs = {}
for r in tbl[1:]:
    segs[r[2]] = {'num': int(r[0]), 'name': r[1], 'id': r[2], 'mustStandards': split_list(r[3]), 'pack': r[4], 'tiers': r[5]}

# Segment details
i = 0
cur = None
while i < len(pd):
    k, lvl, v = pd[i]
    if k == 'p':
        m = re.match(r'V-(\d+) (.+) \((\w+)\)$', v)
        if m:
            cur = m.group(3)
    elif k == 't' and cur and cur in segs:
        rows = {r[0]: r[1] for r in v if len(r) >= 2}
        if 'Activated Macro Processes' in rows:
            s = segs[cur]
            s['activationText'] = rows.get('Activated Macro Processes', '')
            s['kpisText'] = rows.get('KPIs', '')
            s['risks'] = split_list(rows.get('Risks', '')) if 'Same as' not in rows.get('Risks', '') else None
            s['audits'] = split_list(rows.get('Audits', '')) if 'Same as' not in rows.get('Audits', '') else None
            kp = rows.get('KPIs', '')
            kpis = []
            if 'Same as' not in kp:
                for part in re.split(r'\n| / ', kp):
                    mm = re.match(r'([A-Z]+-KPI-\d+)\s+(.*)$', part.strip())
                    if mm:
                        rest = mm.group(2)
                        m2 = re.match(r'(.*?)\s*([<>]=?\s*[\d.,]+\s*%?|<\s*Target|<\s*ISO limit|<\s*\d+ days|[\d.]+%|0|100%)\s*(Monthly|Quarterly|Per Submission)?$', rest)
                        if m2:
                            kpis.append({'code': mm.group(1), 'name': m2.group(1).strip(), 'target': m2.group(2).strip(), 'frequency': (m2.group(3) or 'Monthly')})
                        else:
                            kpis.append({'code': mm.group(1), 'name': rest, 'target': '', 'frequency': 'Monthly'})
            s['kpis'] = kpis if kpis else None
            cur = None
    i += 1
# "Same as PHM" for PHA
for s in segs.values():
    if s.get('kpis') is None:
        if s["id"] != "PHM": s["kpis"] = [dict(k, code=k['code'].replace('PHM', s['id'])) for k in segs['PHM']['kpis']]
    if s.get('risks') is None:
        s['risks'] = segs['PHM']['risks']
    if s.get('audits') is None:
        s['audits'] = segs['PHM']['audits']
cat['segments'] = list(segs.values())

# Macro processes SIPOC
mps = {}
for idx, (k, lvl, v) in enumerate(pd):
    if k == 'p':
        m = re.match(r'MP-(\d{3})\s+·\s+(\w+)\s+—\s+(.+)$', v)
        if m:
            mid = 'MP-' + m.group(1)
            goal = pd[idx + 1][2].replace('Goal: ', '') if pd[idx + 1][0] == 'p' else ''
            t = next(pd[j][2] for j in range(idx + 1, idx + 4) if pd[j][0] == 't')
            rows = {r[0]: r[1] for r in t}
            steps = [re.sub(r'^\d+\.\s*', '', x).strip() for x in re.split(r'\n| / ', rows['Process (P)'])]
            mps[mid] = {'id': mid, 'code': m.group(2), 'name': m.group(3).strip(), 'goal': goal,
                        'sipoc': {'S': split_list(rows['Suppliers (S)']), 'I': split_list(rows['Inputs (I)']),
                                  'P': steps, 'O': split_list(rows['Outputs (O)']), 'C': split_list(rows['Customers (C)'])}}

# Chain relationships
tbl, _ = find_table_after(lambda v: v.startswith('4.2 Chain Relationship Matrix'))
cat['chains'] = [{'from': r[0], 'type': r[1], 'to': [x.strip() for x in r[2].split(',')], 'steps': int(r[3])} for r in tbl[1:]]
tbl, _ = find_table_after(lambda v: v.startswith('4.1 Relationship Types'))
cat['relationTypes'] = [{'type': r[0], 'definition': r[1]} for r in tbl[1:]]

# UF steps
tbl, _ = find_table_after(lambda v: v.startswith('Part 5 — User-Facing Flow'))
cat['ufSteps'] = [{'id': r[0], 'name': r[1], 'mpCodes': r[2], 'e2e': r[3]} for r in tbl[1:]]

# Part 6 detailed E2E specs
det = {}
cur = None
mode = None
for idx, (k, lvl, v) in enumerate(pd):
    if k == 'p':
        m = re.match(r'6\.\d+ (E2E-\d+) — ', v)
        if m:
            cur = m.group(1); det[cur] = {'racsi': [], 'flow': []}; mode = 'head'; continue
        if v.startswith('Part 7'):
            cur = None
        if cur and v == 'RACSI Matrix':
            mode = 'racsi'
        elif cur and v.startswith('Process Flow'):
            mode = 'flow'
    elif k == 't' and cur:
        if mode == 'head':
            rows = {r[0]: r[1] for r in v if len(r) >= 2}
            det[cur].update({'trigger': rows.get('Trigger', ''), 'terminal': rows.get('Terminal State', ''),
                             'ufTasks': rows.get('User-Facing Tasks', ''), 'description': rows.get('Detailed Description', ''),
                             'relations': rows.get('Relationship to Other Chains', ''), 'modules': rows.get('Related Modules', ''),
                             'mpCodes': rows.get('Related Macro Processes', '')})
        elif mode == 'racsi':
            det[cur]['racsi'] = [dict(zip(['activity', 'R', 'A', 'C', 'S', 'I'], r)) for r in v[1:]]
        elif mode == 'flow':
            keys = ['ufId', 'ufName', 'goals', 'inputSuppliers', 'inputs', 'seqTaskIds', 'seqTaskName', 'stepId', 'stepName', 'outputs', 'outputCustomers', 'mpNote', 'bpmn']
            det[cur]['flow'] = [dict(zip(keys, r)) for r in v[1:]]
# narratives
cur = None
in_add5 = False
for k, lvl, v in pd:
    if k == 'p' and v.startswith('Addition 5 — Detailed E2E Narrative'):
        in_add5 = True; continue
    if in_add5 and k == 'p':
        m = re.match(r'(E2E-\d+) — ', v)
        if m:
            cur = m.group(1); continue
        if v.startswith('Final Completion Status'):
            break
        if cur:
            parts = v.split('Business Value: ')
            det[cur]['narrative'] = parts[0].strip()
            det[cur]['businessValue'] = parts[1].strip() if len(parts) > 1 else ''
for e in e2e:
    e.update(det.get(e['id'], {}))

# Standards mapping and packs (Part 7/8)
tbl, _ = find_table_after(lambda v: v.startswith('Part 7 — Standards-to-Macro-Process'))
cat['standardsMap'] = [{'standard': r[0], 'segments': [x.strip() for x in r[1].split(',')], 'mpText': r[2]} for r in tbl[1:]]
tbl, _ = find_table_after(lambda v: v.startswith('Part 8 — Compliance Pack'))
cat['compliancePacks'] = [{'id': r[0], 'name': r[1], 'segment': r[2], 'standards': split_list(r[3])} for r in tbl[1:]]

# Activation matrix (Addition 4)
act = {}
in_add4 = False
for k, lvl, v in pd:
    if k == 'p' and v.startswith('Addition 4'):
        in_add4 = True
    if k == 'p' and v.startswith('Addition 5'):
        in_add4 = False
    if in_add4 and k == 't':
        hdr = v[0]
        for r in v[1:]:
            act[r[0]] = {hdr[i]: r[i] for i in range(1, len(hdr))}
cat['activation'] = act

# Tier 6 RACSI sample and DMS048 BPMN sample
cat['tier6Racsi'] = []
for idx, (k, lvl, v) in enumerate(pd):
    if k == 'p' and v.startswith('RACSI Role Assignments per Tier 6'):
        t = next(pd[j][2] for j in range(idx + 1, idx + 3) if pd[j][0] == 't')
        hdr = t[0]
        for r in t[1:]:
            cat['tier6Racsi'].append({'mp': r[0], 'role': r[1], 'bySegment': {hdr[i]: r[i] for i in range(2, len(hdr)) if r[i] != '—'}})
    if k == 'p' and v.startswith('Sample — DMS048'):
        t = next(pd[j][2] for j in range(idx + 1, idx + 3) if pd[j][0] == 't')
        keys = ['ufId', 'ufName', 'goals', 'inputSuppliers', 'inputs', 'taskId', 'taskName', 'stepId', 'stepName', 'outputs', 'outputCustomers', 'mp', 'bpmn', 'segment']
        cat['dms048Sample'] = [dict(zip(keys, r)) for r in t[1:] if r[0] != '...']
    if k == 'p' and v.startswith('E2E-01 — Context To Strategy — Expanded BPMN'):
        t = next(pd[j][2] for j in range(idx + 1, idx + 3) if pd[j][0] == 't')
        keys = ['ufId', 'ufName', 'goals', 'inputSuppliers', 'inputs', 'taskId', 'taskName', 'stepId', 'stepName', 'outputs', 'outputCustomers', 'mp', 'bpmn']
        cat['e2e01Expanded'] = [dict(zip(keys, r)) for r in t[1:]]

# ---------------------------------------------------------------- Workbook
wb = openpyxl.load_workbook(os.path.join(SRC, 'DynamicMS_PDD_Deliverables_D01D10_D15_D26.xlsx'), read_only=True, data_only=True)

def sheet(name):
    ws = wb[name]
    rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
    hdr = [str(h) for h in rows[0] if h is not None]
    out = []
    for r in rows[1:]:
        out.append({hdr[i]: ('' if r[i] is None else (str(r[i]) if not isinstance(r[i], (int, float)) else r[i])) for i in range(len(hdr))})
    return out

xl_map = {
    'D00a End-to-End Processes': 'e2eWorkbook', 'D00b Functions': 'functions', 'D00c Process Levels': 'processLevels',
    'D00d Sample Process Elements': 'sampleElements', 'D01 Macro Processes': 'mpWorkbook', 'D01a Macro Process Sheets': 'mpSheets',
    'D01b Sheet Interactions': 'sheetInteractions', 'D01c Sheet Tasks': 'sheetTasks', 'D01d Sheet KPIs': 'sheetKpis',
    'D01e Sheet Custom Sections': 'sheetSections', 'D01f Composite Sheets': 'compositeSheets', 'D01g Sheet Monitoring': 'sheetMonitoring',
    'D01h SIPOC Register': 'sipocRegister', 'D02 Tasks & Steps': 'steps', 'D02a Task Procedures': 'procedures',
    'D02b Procedure Steps': 'procedureSteps', 'D03 Business Rules': 'businessRules', 'D03a Actions Registry': 'actions',
    'D04 Controls': 'controls', 'D05 Risks': 'risks', 'D06 KPIs': 'kpis', 'D07 Alerts': 'alerts',
    'D08 Reports & Cockpits': 'reports', 'D08a Doc Templates & Formats': 'docTemplates', 'D08b Document Versions': 'docVersions',
    'D08c MS Policies': 'policies', 'D09 Information Class Model': 'classes', 'D09a Tenant Reference Registry': 'referenceRegistry',
    'D10 Data Dictionary': 'dataDictionary', 'D10a Value Lists': 'valueLists', 'D15 AI Use Cases (Extended)': 'aiUseCases',
    'D15b Role Menus': 'roleMenus', 'D26 Modules & Tiers': 'modules',
}
for sname, key in xl_map.items():
    cat[key] = sheet(sname)

# Merge workbook MP data
for r in cat['mpWorkbook']:
    m = mps[r['Macro_Process_ID']]
    m.update({'e2e': r['E2E_ID'], 'function': r['Function_ID'], 'tier': r['Tier'], 'objective': r['Objective'],
              'trigger': r['Trigger'], 'terminal': r['Terminal_State'], 'ownerRole': r['Owner_Role'], 'module': r['Associated_Module']})
cat['macroProcesses'] = [mps[k] for k in sorted(mps)]
del cat['mpWorkbook']

# ---------------------------------------------------------------- Packs catalog
pk = blocks(os.path.join(SRC, '02_DynamicMS_Packs_Integrations_AddOns_v1.docx'))
def ptable(title):
    for i, (k, l, v) in enumerate(pk):
        if k == 'p' and v.startswith(title):
            for j in range(i + 1, len(pk)):
                if pk[j][0] == 't':
                    return pk[j][2]
tbl = ptable('2.1 Pack Overview Table')
packs = {r[0]: {'id': r[0], 'name': r[1], 'type': r[2], 'target': r[3], 'price': int(re.sub(r'[^\d]', '', r[4]))} for r in tbl[1:]}
for i, (k, l, v) in enumerate(pk):
    if k == 'p':
        m = re.match(r'Pack (DMS-\w+) — ', v)
        if m and m.group(1) in packs:
            t = next(pk[j][2] for j in range(i + 1, i + 3) if pk[j][0] == 't')
            rows = {r[0]: r[1] for r in t}
            packs[m.group(1)].update({'valueProposition': rows.get('Value Proposition', ''), 'modules': rows.get('Related Modules', ''),
                                      'mpText': rows.get('Related Macro Processes', ''), 'deployment': rows.get('Deployment Modes', ''),
                                      'users': rows.get('User Range', ''), 'storage': rows.get('Storage', ''), 'addOnTo': rows.get('Add-On To', '')})
cat['packs'] = list(packs.values())
tbl = ptable('1.4 Deployment Modes Overview')
cat['deploymentModes'] = [{'name': r[0], 'code': r[1], 'description': r[2], 'typical': r[3]} for r in tbl[1:]]
tbl = ptable('5.2 Pricing Multipliers')
mult = {r[1]: float(r[2].replace('x', '')) for r in tbl[1:]}
for d in cat['deploymentModes']:
    d['multiplier'] = mult[d['code']]
tbl = ptable('5.1 The Seven Deployment Modes')
for r in tbl[1:]:
    for d in cat['deploymentModes']:
        if d['code'] == r[1]:
            d.update({'isolation': r[3], 'residency': r[4]})
tbl = ptable('5.3 Full Pricing Matrix')
cat['priceMatrix'] = {r[0]: {tbl[0][i]: int(re.sub(r'[^\d]', '', r[i])) for i in range(1, len(tbl[0]))} for r in tbl[1:]}
tbl = ptable('2.5 Segment-to-Pack Mapping')
cat['segmentPacks'] = [{'segment': r[0].split()[-1], 'base': r[1], 'industry': r[2], 'addons': [x.strip() for x in r[3].split(',')]} for r in tbl[1:]]
tbl = ptable('3.1 Integration Overview Table')
ints = {r[0]: {'id': r[0], 'name': r[1], 'category': r[2]} for r in tbl[1:]}
for i, (k, l, v) in enumerate(pk):
    if k == 'p':
        m = re.match(r'(INT-\d+) — ', v)
        if m:
            t = next(pk[j][2] for j in range(i + 1, i + 3) if pk[j][0] == 't')
            rows = {r[0]: r[1] for r in t}
            ints[m.group(1)].update({'goals': rows.get('Goals', ''), 'examples': rows.get('Example Applications', ''),
                                     'valueProposition': rows.get('Value Proposition', ''), 'packs': [x.strip() for x in rows.get('Related Packs', '').split(',')]})
cat['integrations'] = list(ints.values())
tbl = ptable('4.1 Add-On Overview Table')
adds = {r[0]: {'id': r[0], 'name': r[1], 'category': r[2], 'packsText': r[3], 'price': int(re.sub(r'[^\d]', '', r[4]))} for r in tbl[1:]}
for i, (k, l, v) in enumerate(pk):
    if k == 'p':
        m = re.match(r'(ADD-\d+) — ', v)
        if m:
            t = next(pk[j][2] for j in range(i + 1, i + 3) if pk[j][0] == 't')
            rows = {r[0]: r[1] for r in t}
            adds[m.group(1)].update({'goals': rows.get('Goals', ''), 'valueProposition': rows.get('Value Proposition', '')})
cat['addons'] = list(adds.values())
tbl = ptable('6.1 Recommended Bundles')
cat['bundles'] = [{'segment': r[0], 'bundle': r[1], 'price': int(re.sub(r'[^\d]', '', r[2]))} for r in tbl[1:]]
tbl = ptable('6.2 Bundle Discounts')
cat['bundleDiscounts'] = [{'type': r[0], 'discount': r[1]} for r in tbl[1:]]
tbl = ptable('6.3 Enterprise Agreement')
cat['agreements'] = [dict(zip(['tier', 'users', 'commitment', 'discount', 'sla'], r)) for r in tbl[1:]]
clusters = []
for i, (k, l, v) in enumerate(pk):
    if k == 'p' and re.search(r'Cluster \(', v):
        t = next(pk[j][2] for j in range(i + 1, i + 3) if pk[j][0] == 't')
        rows = {r[0]: r[1] for r in t}
        clusters.append({'name': v, 'behavior': rows.get('Behavior'), 'pain': rows.get('Pain Points'), 'hopes': rows.get('Hopes'), 'fit': rows.get('Pack Fit')})
cat['clusters'] = clusters

# ---------------------------------------------------------------- SRS requirement IDs
srs = blocks(os.path.join(SRC, 'Dynamic_Apps_Standard_SRS_v1.3.docx'))
reqs = []
section = ''
for k, l, v in srs:
    if k == 'p':
        if re.match(r'^\d+\.\d+ ', v):
            section = v
        m = re.match(r'^((?:N?FR)-DA-[A-Z0-9]+-\d+): (.*)$', v)
        if m:
            reqs.append({'id': m.group(1), 'section': section, 'text': m.group(2)})
cat['requirements'] = reqs

# ---------------------------------------------------------------- D30 licensing
lic = blocks(os.path.join(SRC, 'CD_D30_Licensing_Implementation_Schema.docx'))
t = next(v for k, l, v in lic if k == 't')
cat['licensingLinks'] = [dict(zip(['artifact', 'linkId', 'description'], r)) for r in t[1:]]

with open(os.path.join(OUT, 'catalog.en.json'), 'w', encoding='utf-8') as f:
    json.dump(cat, f, ensure_ascii=False, indent=1)
print('macro processes', len(cat['macroProcesses']), 'steps', len(cat['steps']), 'segments', len(cat['segments']),
      'requirements', len(reqs), 'activation', len(act))
